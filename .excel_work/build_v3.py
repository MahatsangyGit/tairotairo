# -*- coding: utf-8 -*-
"""Refonte : offre/capacité, acquisition→croissance, Ampio+ fusionné, indexation, charges patronales, Ask, LTV bornée."""
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as CL

PATH = '/Users/macbookair/Desktop/Etude Tairo/Excel base/Modele_Tairo_3_etats.xlsx'
OUT = '/Users/macbookair/Desktop/tairotairo/.excel_work/Modele_Tairo_3_etats_V3.xlsx'

wb = openpyxl.load_workbook(PATH)

# ---------------------------------------------------------------- styles
C_DARK, C_SEC, C_HDRBG, C_INPUT, C_GRAY = '1F4E79', '4472C4', 'D9E1F2', 'FFF6CC', '595959'
F_TITLE = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
F_SEC = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
F_HDR = Font(name='Calibri', size=10, bold=True, color=C_DARK)
FILL_T = PatternFill('solid', fgColor=C_DARK); FILL_S = PatternFill('solid', fgColor=C_SEC)
FILL_H = PatternFill('solid', fgColor=C_HDRBG); FILL_I = PatternFill('solid', fgColor=C_INPUT)
B_THIN = Border(bottom=Side(style='thin', color='B2B2B2')); B_TOP = Border(top=Side(style='thin', color='404040'))
MONEY, PCT, PCT2, CNT, CNT1, RATIO, INT = '#,##0;(#,##0)', '0.0%', '0.00%', '#,##0', '#,##0.0', '0.00', '0'

def fcalc(v, bold=False, size=10):
    color = '008000' if (isinstance(v, str) and v.startswith('=') and '!' in v) else '000000'
    return Font(name='Calibri', size=size, bold=bold, color=color)

def clear_sheet(ws, rows=200, cols=40):
    for r in range(1, max(ws.max_row, rows) + 1):
        for c in range(1, max(ws.max_column, cols) + 1):
            cell = ws.cell(row=r, column=c)
            cell.value = None; cell.font = Font(name='Calibri', size=10); cell.fill = PatternFill()
            cell.border = Border(); cell.number_format = 'General'; cell.alignment = Alignment()
        ws.row_dimensions[r].height = None

def title(ws, text, lastc):
    ws.cell(row=1, column=1, value=text)
    for c in range(1, lastc + 1):
        cl = ws.cell(row=1, column=c); cl.font = F_TITLE; cl.fill = FILL_T
    ws.row_dimensions[1].height = 22

def section(ws, r, text, lastc):
    ws.cell(row=r, column=1, value=text)
    for c in range(1, lastc + 1):
        cl = ws.cell(row=r, column=c); cl.font = F_SEC; cl.fill = FILL_S
    ws.row_dimensions[r].height = 18

def hdr_row(ws, r, values, lastc):
    for i, v in enumerate(values):
        if v is not None:
            ws.cell(row=r, column=i + 1, value=v)
    for c in range(1, lastc + 1):
        cl = ws.cell(row=r, column=c); cl.font = F_HDR; cl.fill = FILL_H; cl.border = B_THIN

# ---------------------------------------------------------------- PARAMETRES : layout nommé
STD = ['Libellé', 'Valeur base', 'Unité', 'Valeur scénario', 'Note']
SC_VOL, SC_GR, SC_PX, SC_CH, SC_OPX = 'Scenarios!$E$4', 'Scenarios!$E$5', 'Scenarios!$E$6', 'Scenarios!$E$7', 'Scenarios!$E$8'
STRESS_OPX = '*(1+IF({B:st_on}=1,{B:st_opex},0))'

def R(name, label, b, unit, d, note=None, fmt=MONEY):
    return ('row', dict(name=name, label=label, b=b, unit=unit, d=d, note=note, fmt=fmt))

LAYOUT = [
    ('title', 'TAIRO — HYPOTHÈSES (MGA HT) — MODÈLE 3 ÉTATS'),
    ('scen',), ('scen_name',), ('blank',),

    ('section', '1 · REVENUS — PRESTATIONS (Tairo Ampio, cœur du modèle — commission prélevée sur le prestataire)'), ('cols', STD),
    R('presta_m1', 'Prestations demandées au mois 1', 75, 'prestations/mois', '={self}*' + SC_VOL, 'Demande brute. Les prestations réalisées = MIN(demande, capacité de l\'offre) — voir section 10.', CNT),
    R('panier', "Panier moyen d'une prestation (année 1)", 35000, 'Ar', '={self}*' + SC_PX, 'Indexé chaque année sur l\'inflation des coûts variables (section 6).'),
    R('mix10', 'Part du GMV — catégories à 10 % (Plomberie, Ménage, BTP…)', 0.7, '% GMV', '={self}', None, PCT),
    R('mix15', 'Part du GMV — catégories à 15 % (Cuisine)', 0.1, '% GMV', '={self}', None, PCT),
    R('mix20', 'Part du GMV — catégories à 20 % (Déménagement)', 0.05, '% GMV', '={self}', None, PCT),
    R('mix25', 'Part du GMV — catégories à 25 % (Informatique, Transport…)', 0.15, '% GMV', '={self}', None, PCT),
    R('take_presta', 'Take rate moyen services (calculé)', None, '%', '=0.1*{mix10}+0.15*{mix15}+0.2*{mix20}+0.25*{mix25}', None, PCT2),
    R('ctrl_mix', 'Contrôle mix (doit = 100 %)', None, '%', '={mix10}+{mix15}+{mix20}+{mix25}', None, PCT),
    ('blank',),
    ('section', "2 · REVENUS — ABONNEMENT Ampio+ (fusionné avec l'offre : part des prestataires actifs abonnés)"), ('cols', STD),
    R('prix_abo', "Prix de l'abonnement mensuel (année 1)", 24000, 'Ar/mois', '={self}', 'Indexé chaque année (section 6).'),
    R('pen_abo', 'Part des prestataires actifs abonnés à Ampio+ (pénétration)', 0.10, '%', '={self}', "Remplace l'ancien stock d'abonnés autonome (5 abonnés, churn 6 %) : les abonnés suivent désormais le nombre de prestataires actifs (section 10).", PCT),
    ('blank',),
    ('section', '3 · REVENUS — MISE EN LIGNE VIDÉO (Tairo Ampianaro B2B, clients professionnels)'), ('cols', STD),
    R('prix_video', 'Prix mensuel du service vidéo (année 1)', 64000, 'Ar/mois', '={self}', 'Indexé chaque année (section 6).'),
    R('video_m1', 'Clients vidéo actifs au mois 1', 15, 'clients', '={self}*' + SC_VOL, None, CNT),
    R('churn_video', 'Churn mensuel clients vidéo', 0.025, '%/mois', '=MIN(1,{self}*' + SC_CH + ')', None, PCT),
    ('blank',),
    ('section', '4 · REVENUS — LOCATION MATÉRIEL PRO (Ampindramo PRO, commission par paliers)'), ('cols', STD),
    R('locpro_m1', 'Locations pro demandées au mois 1', 15, 'locations/mois', '={self}*' + SC_VOL, None, CNT),
    R('loyer_pro', "Loyer total moyen d'une location pro (année 1)", 15000, 'Ar', '={self}*' + SC_PX),
    R('pal30', 'Part du GMV — total ≤ 25 000 Ar (taux 30 %)', 0.85, '% GMV', '={self}', None, PCT),
    R('pal20', 'Part du GMV — 25 001 à 45 000 Ar (taux 20 %)', 0.12, '% GMV', '={self}', None, PCT),
    R('pal15', 'Part du GMV — > 45 000 Ar (taux 15 %)', 0.03, '% GMV', '={self}', None, PCT),
    R('take_locpro', 'Take rate moyen location pro (calculé)', None, '%', '=0.3*{pal30}+0.2*{pal20}+0.15*{pal15}', None, PCT2),
    R('ctrl_pal', 'Contrôle mix (doit = 100 %)', None, '%', '={pal30}+{pal20}+{pal15}', None, PCT),
    ('blank',),
    ('section', '5 · REVENUS — LOCATION MATÉRIEL PARTICULIERS (Ampindramo C2C)'), ('cols', STD),
    R('c2c_m1', 'Locations particuliers demandées au mois 1', 12, 'locations/mois', '={self}*' + SC_VOL, None, CNT),
    R('loyer_c2c', 'Loyer total moyen (année 1)', 10000, 'Ar', '={self}*' + SC_PX),
    R('take_c2c', 'Taux de commission particuliers (fixe, hors caution)', 0.13, '%', '={self}', None, PCT),
    ('blank',),
    ('section', '6 · INDEXATION DES PRIX & INFLATION (pas annuel, appliqué à chaque nouvelle année du plan)'), ('cols', STD),
    R('index_prix', 'Indexation annuelle des prix (abonnement Ampio+, service vidéo)', 0.08, '%/an', '={self}', 'Aligné par défaut sur l\'inflation des coûts pour éviter la compression silencieuse des marges.', PCT),
    R('infl_var', 'Inflation annuelle des coûts variables et des paniers / loyers moyens', 0.08, '%/an', '={self}', None, PCT),
    R('infl_opex', 'Inflation annuelle des OPEX (salaires, loyer, tech, accompagnement)', 0.08, '%/an', '={self}', None, PCT),
    ('blank',),
    ('section', "7 · CROISSANCE MENSUELLE DE LA DEMANDE — PAR ANNÉE, AU BUDGET D'ACQUISITION DE RÉFÉRENCE (ralentissement après A1)"),
    ('cols', ['Segment', 'Année 1 (base)', 'Année 2 (base)', 'Année 3 (base)', 'Note', 'Année 1 (actif)', 'Année 2 (actif)', 'Année 3 (actif)']),
    ('growth', dict(name='g_video', label='Clients vidéo B2B', vals=(0.05, 0.04, 0.03),
                    note='Ralentissement introduit : 5 % → 4 % → 3 %. Multiplié par le facteur d\'acquisition (section 9).')),
    ('growth', dict(name='g_presta', label='Prestations Ampio (demande)', vals=(0.02, 0.025, 0.02),
                    note="2 % pré-PMF, 2,5 % en A2 si collab 2 actif (sinon plafond section 8), 2 % en A3. L'accélération à 3 % a été retirée.")),
    ('growth', dict(name='g_locpro', label='Locations pro (demande)', vals=(0.018, 0.025, 0.02), note='Même logique de plafond que les prestations tant que collab 2 absent.')),
    ('growth', dict(name='g_c2c', label='Locations particuliers C2C (demande)', vals=(0.015, 0.012, 0.01), note='C2C plus lent (confiance, caution, concurrence WhatsApp/Facebook). Pas de bonus d\'embauche.')),
    ('blank',),
    ('section', '8 · PLAFOND DE CROISSANCE SI COLLAB 2 ABSENT (années 2–3)'), ('cols', STD),
    R('cap_presta', 'Prestations Ampio — MoM tant que collab 2 inactif', 0.02, '%/mois', '={self}*' + SC_GR, "Sans 3e personne, l'onboarding n'accélère pas. C2C non plafonné.", PCT),
    R('cap_locpro', 'Locations pro — MoM tant que collab 2 inactif', 0.018, '%/mois', '={self}*' + SC_GR, None, PCT),
    ('blank',),
    ('section', "9 · ACQUISITION — BUDGET PUBLICITÉ PAR ANNÉE ET LIEN AVEC LA CROISSANCE (élasticité)"), ('cols', STD),
    R('pub_a1', 'Budget publicité / acquisition — année 1', 400000, 'Ar/mois', '={self}', "Remplace la règle « 6 % du CA » et le plancher cold-start. Saisi par année, non indexé."),
    R('pub_a2', 'Budget publicité / acquisition — année 2', 500000, 'Ar/mois', '={self}'),
    R('pub_a3', 'Budget publicité / acquisition — année 3', 600000, 'Ar/mois', '={self}'),
    R('pub_ref', 'Budget de référence (celui auquel la table de croissance section 7 est calibrée)', 400000, 'Ar/mois', '={self}'),
    R('elast', 'Élasticité de la croissance au budget (croissance × (budget M-1 / référence)^élasticité)', 0.5, 'coef.', '={self}', "0,5 = rendements décroissants : doubler le budget augmente la croissance de 41 %. Budget du mois précédent (pas de circularité).", RATIO),
    R('elast_min', 'Borne basse du multiplicateur de croissance', 0.5, 'coef.', '={self}', None, RATIO),
    R('elast_max', 'Borne haute du multiplicateur de croissance', 1.5, 'coef.', '={self}', None, RATIO),
    R('acq_share', "Part de l'accompagnement marketing (OPEX) dédiée à l'acquisition (CAC)", 0.8, '%', '={self}', None, PCT),
    R('founder_time', "Part du temps fondateur dédié à l'acquisition (valorisée dans le CAC)", 0.3, '%', '={self}', None, PCT),
    R('new_share', 'Part des transactions réalisées issues de nouveaux clients (demande)', 0.15, '%', '={self}', "Hypothèse à valider avec les données d'usage.", PCT),
    R('tx_per_new', 'Nombre moyen de transactions par nouveau client (par mois)', 1, 'transactions/client/mois', '={self}', None, RATIO),
    R('life_tx', "Durée de vie moyenne d'un client transactionnel (mois)", 2, 'mois', '={self}', 'Clients ponctuels de la marketplace, moins engagés qu\'un abonnement.', INT),
    R('split_dem', "Part du budget d'acquisition allouée à la DEMANDE", 0.7, '%', '={self}', None, PCT),
    R('split_off', "Part du budget d'acquisition allouée à l'OFFRE", 0.3, '%', '={self}', 'Standard marketplace : acquérir la demande coûte plus cher.', PCT),
    R('ctrl_split', 'Contrôle split (doit = 100 %)', None, '%', '={split_dem}+{split_off}', None, PCT),
    ('blank',),
    ('section', "10 · OFFRE — PRESTATAIRES, LOUEURS, ANNONCES (onboarding, churn, capacité CONTRAIGNANTE)"), ('cols', STD),
    R('prest_m1', 'Prestataires Ampio actifs au mois 1', 25, 'prestataires', '={self}*' + SC_VOL, None, CNT),
    R('prest_onb', 'Prestataires onboardés par mois — sans collab 2', 3, 'prestataires/mois', '={self}*' + SC_GR, 'Capacité d\'onboarding du fondateur + collab 1.', CNT1),
    R('prest_onb2', 'Prestataires onboardés par mois — avec collab 2', 8, 'prestataires/mois', '={self}*' + SC_GR, 'Le collab 2 est dédié au support / onboarding.', CNT1),
    R('prest_churn', 'Churn mensuel des prestataires', 0.03, '%/mois', '=MIN(1,{self}*' + SC_CH + ')', None, PCT),
    R('jobs_max', 'Prestations max par prestataire et par mois', 8, 'jobs/pers/mois', '={self}', None, CNT),
    R('loueurs_m1', 'Loueurs pro actifs au mois 1', 8, 'loueurs', '={self}*' + SC_VOL, None, CNT),
    R('loueurs_onb', 'Loueurs pro onboardés par mois — sans collab 2', 1, 'loueurs/mois', '={self}*' + SC_GR, None, CNT1),
    R('loueurs_onb2', 'Loueurs pro onboardés par mois — avec collab 2', 2, 'loueurs/mois', '={self}*' + SC_GR, None, CNT1),
    R('loueurs_churn', 'Churn mensuel des loueurs pro', 0.02, '%/mois', '=MIN(1,{self}*' + SC_CH + ')', None, PCT),
    R('loc_max', 'Locations max par loueur et par mois', 4, 'loc/mois', '={self}', None, CNT),
    R('annonces_m1', 'Annonces C2C actives au mois 1', 15, 'annonces', '={self}*' + SC_VOL, None, CNT),
    R('annonces_onb', 'Nouvelles annonces C2C par mois', 2, 'annonces/mois', '={self}*' + SC_GR, 'Non lié au collab 2 (peu d\'onboarding).', CNT1),
    R('annonces_churn', 'Churn mensuel des annonces C2C', 0.05, '%/mois', '=MIN(1,{self}*' + SC_CH + ')', None, PCT),
    R('ann_max', 'Locations max par annonce et par mois', 1, 'loc/mois', '={self}', None, CNT),
    R('fill_max', "Taux d'utilisation maximal de la capacité (fill rate max)", 0.8, '%', '={self}', "Volumes réalisés = MIN(demande, offre × max × fill max). Remplace la rampe de fill rate saisie et les KPI de liquidité tautologiques.", PCT),
    ('blank',),
    ('section', "11 · LANCEMENT (M1–M6) — subvention d'offre sur les prestations"), ('cols', STD),
    R('launch_dur', 'Durée de la phase de lancement', 6, 'mois', '={self}', None, INT),
    R('launch_tr', 'Réduction du take rate Ampio pendant le lancement', 0.02, '% (0,02 = 2 points)', '={self}', "Corrigé : était saisi « 2 » (200 points), ce qui annulait tout le CA commissions M1–M6. Mix 13,25 % → 11,25 %.", PCT2),
    ('blank',),
    ('section', '12 · COÛTS VARIABLES — COGS ET COÛTS OPÉRATIONNELS LIÉS AU VOLUME'), ('cols', STD),
    R('sms', 'SMS de notification (2 × 35,71 Ar) — prestations Ampio + locations C2C', 71.42, 'Ar/transaction', '={self}', None, '#,##0.00'),
    R('mm', 'Commission Mobile Money par transaction', 0.03, '%', '={self}', "3 % du montant brut : CA abonnements Ampio+, GMV prestations, GMV locations pro et C2C. Exclut le service vidéo (paiement B2B).", PCT),
    R('support', 'Support / modération par transaction (Ampio + loc pro + C2C)', 1000, 'Ar/transaction', '={self}*' + SC_OPX + STRESS_OPX),
    R('verif', "Vérification / onboarding d'un nouveau client (prestataire, loueur, client vidéo, client demande)", 3000, 'Ar/nouveau', '={self}*' + SC_OPX + STRESS_OPX),
    R('litiges', 'Provision remboursements / litiges', 0.01, '% du GMV', '={self}*' + SC_OPX + STRESS_OPX, None, PCT),
    ('blank',),
    ('section', '13 · OPEX MENSUEL (charges fixes, année 1)'), ('cols', STD),
    R('loyer', 'Loyer et domiciliation', 125000, 'Ar/mois', '={self}*' + SC_OPX + STRESS_OPX, 'Wezara'),
    R('sal_fond', 'Salaire brut fondateur', 800000, 'Ar/mois', '={self}*' + SC_OPX + STRESS_OPX),
    R('sal_c1', 'Salaire brut collaborateur 1', 1000000, 'Ar/mois', '={self}*' + SC_OPX + STRESS_OPX),
    R('charges_taux', 'Taux de charges patronales (CNaPS 13 % + organisation sanitaire 5 %)', 0.18, '% des salaires bruts', '={self}', "Appliqué à tous les salaires (fondateur, collab 1, collab 2).", PCT),
    R('charges_calc', 'Charges patronales fondateur + collab 1 (calculé)', None, 'Ar/mois', '={charges_taux}*({sal_fond}+{sal_c1})'),
    R('heberg', 'Hébergement web', 16583, 'Ar/mois', '={self}*' + SC_OPX + STRESS_OPX, 'Simafri'),
    R('storage', 'Object Storage', 30000, 'Ar/mois', '={self}*' + SC_OPX + STRESS_OPX, 'Cloudflare R2'),
    R('moder', 'Modération', 10000, 'Ar/mois', '={self}*' + SC_OPX + STRESS_OPX, 'Crédit data'),
    R('accomp', 'Marketing — accompagnement (fixe, hors budget publicité section 9)', 100000, 'Ar/mois', '={self}*' + SC_OPX + STRESS_OPX),
    R('total_opex', 'TOTAL OPEX mensuel (année 1, charges patronales incluses, hors publicité et collab 2)', None, 'Ar/mois',
      '={loyer}+{sal_fond}+{sal_c1}+{charges_calc}+{heberg}+{storage}+{moder}+{accomp}', "Base du besoin de démarrage et de la trésorerie de sécurité (Financement)."),
    ('blank',),
    ('section', "14 · CAPACITÉ D'EXÉCUTION — 2E COLLABORATEUR (embauche + déclencheur de trésorerie)"), ('cols', STD),
    R('sal_c2', 'Salaire brut mensuel collab 2 (base, avant inflation ; charges patronales ajoutées au P&L)', 400000, 'Ar/mois', '={self}*' + SC_OPX),
    R('c2_min_month', 'Mois minimum d\'embauche', 13, 'n° de mois (1-36)', '={self}', "L'embauche réelle = 1er mois ≥ ce mois où trésorerie > plancher ET EBITDA > 0 pendant N mois consécutifs (décalage d'un mois). Poste conservé ensuite.", INT),
    R('c2_plancher', 'Plancher de trésorerie exigé pour embaucher (référence section 16)', None, 'Ar', '={fin_plancher}'),
    R('c2_consec', 'Nombre de mois consécutifs exigés', 3, 'mois', '={self}', None, INT),
    R('c2_ebitda', 'Exiger aussi EBITDA mensuel > 0 (1 = oui)', 1, '0/1', '={self}', 'Les deux conditions portent sur les mois strictement antérieurs.', INT),
    ('blank',),
    ('section', '15 · CAPEX (investissements)'),
    ('cols', ['Équipement', 'Quantité', 'PU (Ar TTC)', 'Total (Ar)', 'Durée (ans)', 'Note']),
    ('capex', dict(name='capex1', label='Ordinateur portable (IAS 16 — corporelle)', qty=1, pu=900000, dur=5, note='Ordinateur windows')),
    ('capex', dict(name='capex2', label='Design graphique UI/UX (IAS 38 — incorporelle)', qty=1, pu=4000000, dur=5, note=None)),
    ('capex_total', dict(name='total_capex', label='TOTAL CAPEX')),
    ('blank',),
    ('section', '16 · FINANCEMENT'), ('cols', STD),
    R('apport', 'Apport en fonds propres', 9000000, 'Ar', '={self}', "Apport des fondateurs. Le solde du besoin (Ask) est modélisé en emprunt dans la trésorerie."),
    R('fin_mois_dem', 'Besoin de démarrage — nombre de mois de charges couverts', 12, 'mois', '={self}', 'N mois × (OPEX fixes chargés + budget publicité A1 + coûts variables du mois 1). Calcul dans Financement.', INT),
    R('fin_mois_sec', "Trésorerie de sécurité — nombre de mois d'OPEX fixes", 3, 'mois', '={self}', 'Cible de la validation 3 ans : la trésorerie ne doit jamais passer sous ce matelas.', INT),
    R('fin_plancher', "Plancher de trésorerie pour l'embauche du collaborateur 2", 1000000, 'Ar', '={self}', 'Découplé de la trésorerie de sécurité ; utilisé uniquement par le déclencheur d\'embauche (section 14).'),
    R('taux', "Taux d'emprunt annuel", 0.09, '%/an', '={self}', None, PCT),
    R('duree', "Durée de l'emprunt (mois)", 60, 'mois', '={self}', "Si > 36, le capital restant dû au mois 36 est affiché dans Financement et au Bilan.", INT),
    R('ask_inv', "Ask — part visée auprès d'investisseurs (equity / quasi-fonds propres), indicatif", 0.3, '%', '={self}', "Ventilation de présentation uniquement : la trésorerie modélise l'ensemble du financement externe en emprunt (mécanique inchangée).", PCT),
    R('ref_v23', 'Référence V2.3 — apport calé à la main pour un plancher de 1 M Ar (documentation)', 36700000, 'Ar', '={self}', 'Ancien apport itéré manuellement. Comparé au besoin réel calculé (Financement, bloc Validation 3 ans).'),
    ('blank',),
    ('section', '17 · FISCALITÉ MADAGASCAR (LF 2026) — IS SYNTHÉTIQUE / IR RÉGIME RÉEL'), ('cols', STD),
    R('fisc_seuil', 'Seuil de passage au régime réel (CA annuel)', 200000000, 'Ar/an', '={self}'),
    R('fisc_is', 'Taux Impôt Synthétique si CA < seuil (sur le CA)', 0.05, '% CA', '={self}', None, PCT),
    R('fisc_ir', 'Taux IR si CA ≥ seuil (sur le bénéfice)', 0.2, '% bénéfice', '={self}', None, PCT),
    R('fisc_min_fixe', 'Minimum de perception IR — part fixe (régime réel uniquement)', 1000000, 'Ar/an', '={self}', "Code des impôts : 1 000 000 Ar pour les « autres entreprises ». Inactif tant que CA < seuil."),
    R('fisc_min_pct', 'Minimum de perception IR — % du CA (régime réel uniquement)', 0.01, '% CA', '={self}', None, PCT),
    ('blank',),
    ('section', '18 · STRESS TEST — SCÉNARIO DÉGRADÉ COMBINÉ (indépendant de Prudent / Base / Optimiste)'),
    ('cols', ['Libellé', 'Valeur', 'Unité', None, 'Note']),
    R('st_on', 'Activer le stress test (1 = oui, 0 = non)', 0, '0/1', None, 'Interrupteur unique : bascule sur 1 pour appliquer tous les chocs ci-dessous.', INT),
    R('st_demande', 'Choc de demande / volumes de départ (réduction)', 0.3, '%', None, 'Appliqué aux volumes et à l\'offre du mois 1 de chaque segment.', PCT),
    R('st_retard', 'Retard de lancement (croissance de la demande gelée)', 6, 'mois', None, None, INT),
    R('st_opex', 'Choc OPEX (hausse) — appliqué aux charges fixes et coûts opérationnels', 0.3, '%', None, None, PCT),
    R('st_capex', 'Choc CAPEX (hausse)', 0.2, '%', None, None, PCT),
    R('st_impaye', "Taux d'impayé sur encaissements", 0.04, '%', None, None, PCT),
    R('st_takerate', 'Réduction du take rate (points, 0,025 = 2,5 points)', 0.025, '%', None, None, PCT2),
    R('st_delai', "Délai d'encaissement client", 1, 'mois (≈30 jours)', None, None, INT),
]

# positions
ROW = {}       # name -> row
placed = []
r = 0
for item in LAYOUT:
    r += 1
    placed.append((r, item))
    if item[0] in ('row', 'growth', 'capex', 'capex_total'):
        ROW[item[1]['name']] = r
    elif item[0] == 'scen':
        ROW['scen'] = r
LASTP = r

def resolve(tpl, local=False, self_row=None):
    def sub(m):
        col, name = m.group(1), m.group(2)
        if name == 'self':
            return f'B{self_row}'
        col = col or 'D'
        ref = f'${col}${ROW[name]}'
        return ref if local else f'Parametres!{ref}'
    return re.sub(r'\{(?:([A-H]):)?([a-z0-9_]+)\}', sub, tpl)

def PR(name, col='D'):
    return f'Parametres!${col}${ROW[name]}'

# ---------------------------------------------------------------- remap des anciennes références Parametres (onglets conservés)
OLD2NAME = {2: 'scen', 3: 'scen_name', 7: 'prix_abo', 13: 'prix_video', 14: 'video_m1', 15: 'churn_video', 19: 'presta_m1', 20: 'panier',
            25: 'take_presta', 30: 'locpro_m1', 31: 'loyer_pro', 35: 'take_locpro', 40: 'c2c_m1', 41: 'loyer_c2c', 42: 'take_c2c',
            59: 'sms', 60: 'mm', 61: 'support', 62: 'verif', 63: 'litiges', 64: 'infl_var', 68: 'loyer', 69: 'sal_fond', 70: 'sal_c1',
            71: 'heberg', 72: 'storage', 73: 'moder', 74: 'accomp', 75: 'total_opex', 76: 'infl_opex', 80: 'acq_share', 81: 'new_share',
            82: 'tx_per_new', 83: 'founder_time', 85: 'life_tx', 86: 'split_dem', 87: 'split_off', 92: 'launch_dur', 94: 'launch_tr',
            98: 'sal_c2', 99: 'c2_min_month', 100: 'c2_min_month', 101: 'c2_plancher', 102: 'c2_consec', 103: 'c2_ebitda',
            107: 'capex1', 108: 'capex2', 109: 'total_capex', 113: 'apport', 114: 'fin_mois_dem', 115: 'fin_mois_sec', 116: 'fin_plancher',
            117: 'taux', 118: 'duree', 119: 'ref_v23', 123: 'fisc_seuil', 124: 'fisc_is', 125: 'fisc_ir', 126: 'fisc_min_fixe', 127: 'fisc_min_pct',
            131: 'prest_m1', 132: 'jobs_max', 136: 'loueurs_m1', 137: 'loc_max', 138: 'annonces_m1', 139: 'ann_max',
            143: 'st_on', 144: 'st_demande', 145: 'st_retard', 146: 'st_opex', 147: 'st_capex', 148: 'st_impaye', 149: 'st_takerate', 150: 'st_delai'}
ROW['scen_name'] = 3
KEEP = ['P&L', 'Contribution', 'Tresorerie', 'Amortissements', 'Scenarios', 'Bilan', 'Base graphique']
ppat = re.compile(r"Parametres!(\$?)([A-H])(\$?)(\d{1,3})")
def remap_param(f):
    def sub(m):
        old = int(m.group(4))
        if old not in OLD2NAME:
            raise KeyError(f'ref Parametres ligne {old} non mappée : {f}')
        return f"Parametres!{m.group(1)}{m.group(2)}{m.group(3)}{ROW[OLD2NAME[old]]}"
    return ppat.sub(sub, f)
WRITTEN = set()   # cellules réécrites avec les nouvelles références (à exclure du remap)
def do_remap_keep():
    for sn in KEEP:
        for row in wb[sn].iter_rows():
            for c in row:
                if (sn, c.coordinate) in WRITTEN:
                    continue
                if isinstance(c.value, str) and c.value.startswith('=') and 'Parametres!' in c.value:
                    c.value = remap_param(c.value)

# ---------------------------------------------------------------- écriture de Parametres
P = wb['Parametres']
P.data_validations.dataValidation = []
clear_sheet(P, rows=160, cols=8)
for rr, item in placed:
    k = item[0]
    if k == 'title':
        title(P, item[1], 8)
    elif k == 'scen':
        P.cell(row=rr, column=1, value='Scénario actif (1 = Prudent, 2 = Base, 3 = Optimiste)').font = Font(name='Calibri', size=10, bold=True)
        b = P.cell(row=rr, column=2, value=2); b.font = Font(name='Calibri', size=11, bold=True, color='0000FF'); b.fill = FILL_I
    elif k == 'scen_name':
        P.cell(row=rr, column=1, value='Nom du scénario').font = Font(name='Calibri', size=10)
        b = P.cell(row=rr, column=2, value='=CHOOSE(B2,"Prudent","Base","Optimiste")'); b.font = Font(name='Calibri', size=11, bold=True)
    elif k == 'section':
        section(P, rr, item[1], 8)
    elif k == 'cols':
        hdr_row(P, rr, item[1], 8)
    elif k == 'row':
        d = item[1]
        a = P.cell(row=rr, column=1, value=d['label']); a.font = Font(name='Calibri', size=10); a.alignment = Alignment(wrap_text=True, vertical='top')
        if d['b'] is not None:
            b = P.cell(row=rr, column=2, value=d['b']); b.font = Font(name='Calibri', size=10, color='0000FF'); b.fill = FILL_I; b.number_format = d['fmt']
        u = P.cell(row=rr, column=3, value=d['unit']); u.font = Font(name='Calibri', size=10, color=C_GRAY)
        if d['d']:
            f = resolve(d['d'], local=True, self_row=rr)
            dc = P.cell(row=rr, column=4, value=f); dc.font = fcalc(f); dc.number_format = d['fmt']
        if d['note']:
            n = P.cell(row=rr, column=5, value=d['note']); n.font = Font(name='Calibri', size=9, italic=True, color=C_GRAY); n.alignment = Alignment(wrap_text=True, vertical='top')
        if d['label'].startswith('TOTAL') or 'calculé' in d['label'] or d['label'].startswith('Contrôle'):
            for c in range(1, 9):
                cl = P.cell(row=rr, column=c); cl.font = Font(name='Calibri', size=cl.font.size or 10, bold=True, italic=cl.font.italic, color=cl.font.color); cl.border = B_TOP
    elif k == 'growth':
        d = item[1]
        a = P.cell(row=rr, column=1, value=d['label']); a.font = Font(name='Calibri', size=10); a.alignment = Alignment(wrap_text=True, vertical='top')
        for i, v in enumerate(d['vals']):
            b = P.cell(row=rr, column=2 + i, value=v); b.font = Font(name='Calibri', size=10, color='0000FF'); b.fill = FILL_I; b.number_format = PCT
            fcell = P.cell(row=rr, column=6 + i, value=f'={CL(2 + i)}{rr}*{SC_GR}'); fcell.font = fcalc(fcell.value); fcell.number_format = PCT2
        n = P.cell(row=rr, column=5, value=d['note']); n.font = Font(name='Calibri', size=9, italic=True, color=C_GRAY); n.alignment = Alignment(wrap_text=True, vertical='top')
    elif k == 'capex':
        d = item[1]
        P.cell(row=rr, column=1, value=d['label']).font = Font(name='Calibri', size=10)
        for col, v, fmt in ((2, d['qty'], INT), (3, d['pu'], MONEY), (5, d['dur'], INT)):
            b = P.cell(row=rr, column=col, value=v); b.font = Font(name='Calibri', size=10, color='0000FF'); b.fill = FILL_I; b.number_format = fmt
        dc = P.cell(row=rr, column=4, value=f'=B{rr}*C{rr}'); dc.font = fcalc(dc.value); dc.number_format = MONEY
        if d['note']:
            P.cell(row=rr, column=6, value=d['note']).font = Font(name='Calibri', size=9, italic=True, color=C_GRAY)
    elif k == 'capex_total':
        P.cell(row=rr, column=1, value=item[1]['label']).font = Font(name='Calibri', size=10, bold=True)
        dc = P.cell(row=rr, column=4, value=f'=SUM(D{ROW["capex1"]}:D{ROW["capex2"]})'); dc.font = Font(name='Calibri', size=10, bold=True); dc.number_format = MONEY
        for c in range(1, 9): P.cell(row=rr, column=c).border = B_TOP
dv = DataValidation(type='list', formula1='"1,2,3"', allow_blank=False, showErrorMessage=True); P.add_data_validation(dv); dv.add('B2')
for col, w in {'A': 66, 'B': 14, 'C': 22, 'D': 15, 'E': 72, 'F': 12, 'G': 12, 'H': 12}.items():
    P.column_dimensions[col].width = w
P.freeze_panes = 'A4'; P.sheet_view.showGridLines = False

# ---------------------------------------------------------------- PLAN VENTES (colonnes B..AK = M1..M36)
PV = wb['Plan ventes']
clear_sheet(PV, rows=80, cols=37)
NCOL = 37
title(PV, "PLAN DE VENTES — acquisition, offre (capacité contraignante), demande, GMV, CA (M1–M36) + flag embauche", NCOL)
hdr_row(PV, 2, ['Année'], NCOL); hdr_row(PV, 3, ['Mois'], NCOL)
for i in range(1, 37):
    L = CL(1 + i)
    PV[f'{L}3'] = i; PV[f'{L}3'].number_format = INT
    PV[f'{L}2'] = f'=INT(({L}3-1)/12)+1'; PV[f'{L}2'].number_format = INT

pv_rows = {}   # key -> row
def pv_line(rw, key, label, f_first, f_next, fmt=MONEY, bold=False):
    """f_first / f_next : fonctions (L, Lp, i) -> formule ; Lp = colonne précédente."""
    pv_rows[key] = rw
    a = PV.cell(row=rw, column=1, value=label); a.font = Font(name='Calibri', size=10, bold=bold)
    for i in range(1, 37):
        L, Lp = CL(1 + i), CL(i)
        f = f_first(L, Lp, i) if i == 1 else f_next(L, Lp, i)
        c = PV.cell(row=rw, column=1 + i, value=f); c.number_format = fmt; c.font = fcalc(f, bold)

def Y(L): return f'{L}$2'
def MO(L): return f'{L}$3'
def infl(name, L): return f'(1+{PR(name)})^({Y(L)}-1)'
STRESS_M1 = f'(1-IF({PR("st_on","B")}=1,{PR("st_demande","B")},0))'
def RETARD(L): return f'AND({PR("st_on","B")}=1,{MO(L)}<={PR("st_retard","B")})'
def GROWTH(name, L): return f'CHOOSE({Y(L)},{PR(name,"F")},{PR(name,"G")},{PR(name,"H")})'

rw = 5
section(PV, rw, 'ACQUISITION → CROISSANCE', NCOL); rw += 1
pv_line(rw, 'pub', 'Budget publicité / acquisition du mois (Ar, saisi par année — section 9)',
        lambda L, Lp, i: f'=CHOOSE({Y(L)},{PR("pub_a1")},{PR("pub_a2")},{PR("pub_a3")})',
        lambda L, Lp, i: f'=CHOOSE({Y(L)},{PR("pub_a1")},{PR("pub_a2")},{PR("pub_a3")})'); rw += 1
pv_line(rw, 'mult', "Multiplicateur de croissance lié au budget (élasticité, budget du mois précédent, borné)",
        lambda L, Lp, i: '=1',
        lambda L, Lp, i: f'=MIN({PR("elast_max")},MAX({PR("elast_min")},IF({PR("pub_ref")}=0,1,({Lp}{pv_rows["pub"]}/{PR("pub_ref")})^{PR("elast")})))', RATIO); rw += 1
pv_line(rw, 'retard', 'Croissance gelée (stress test : retard de lancement) — 1 = oui',
        lambda L, Lp, i: f'=IF({RETARD(L)},1,0)', lambda L, Lp, i: f'=IF({RETARD(L)},1,0)', INT); rw += 2

def supply_block(rw, sec_title, key, m1, onb, onb2, churn, maxk, dem_first_param, gname, capname, price_param, infl_name, take_expr, gmv_label, ca_label, tx_label, with_abo=False):
    """tx_label = (nouveaux..., perdus..., actifs...) ; gmv_label = (minuscule, Capitalisé)."""
    lab_new, lab_lost, lab_stock = tx_label
    gmv_lo, gmv_up = gmv_label
    section(PV, rw, sec_title, NCOL); rw += 1
    FLAG = pv_rows['flag']  # défini plus bas (ligne connue à l'avance)
    pv_line(rw, key + '_new', f'{lab_new} (M1 = stock initial)',
            lambda L, Lp, i: f'={PR(m1)}*{STRESS_M1}',
            (lambda L, Lp, i: f'=IF({Lp}{FLAG}=1,{PR(onb2)},{PR(onb)})') if onb2 else (lambda L, Lp, i: f'={PR(onb)}'), CNT1); rw += 1
    pv_line(rw, key + '_lost', f'{lab_lost} (churn)',
            lambda L, Lp, i: '=0', lambda L, Lp, i: f'={Lp}{rw+1}*{PR(churn)}', CNT1); rw += 1
    pv_line(rw, key + '_stock', f'{lab_stock} (stock × (1 − churn) + nouveaux)',
            lambda L, Lp, i: f'={L}{rw-2}',
            lambda L, Lp, i: f'={Lp}{rw}*(1-{PR(churn)})+{L}{rw-2}', CNT1, bold=True); rw += 1
    pv_line(rw, key + '_cap', "Capacité mensuelle (actifs × max par mois × fill rate max)",
            lambda L, Lp, i: f'={L}{rw-1}*{PR(maxk)}*{PR("fill_max")}', lambda L, Lp, i: f'={L}{rw-1}*{PR(maxk)}*{PR("fill_max")}', CNT1); rw += 1
    if with_abo:
        pv_line(rw, 'abo', 'Abonnés Ampio+ (pénétration × prestataires actifs)',
                lambda L, Lp, i: f'={L}{pv_rows[key+"_stock"]}*{PR("pen_abo")}', lambda L, Lp, i: f'={L}{pv_rows[key+"_stock"]}*{PR("pen_abo")}', CNT1); rw += 1
        pv_line(rw, 'prix_abo', "Prix de l'abonnement indexé", lambda L, Lp, i: f'={PR("prix_abo")}*{infl("index_prix", L)}',
                lambda L, Lp, i: f'={PR("prix_abo")}*{infl("index_prix", L)}'); rw += 1
        pv_line(rw, 'ca_abo', 'CA abonnements Ampio+ (MRR)', lambda L, Lp, i: f'={L}{rw-2}*{L}{rw-1}', lambda L, Lp, i: f'={L}{rw-2}*{L}{rw-1}', bold=True); rw += 1
    # demande
    def dem_next(L, Lp, i):
        g = GROWTH(gname, L)
        if capname:
            g = f'IF(AND({Y(L)}>=2,{Lp}{FLAG}=0),{PR(capname)},{g})'
        return f'={Lp}{pv_rows[key+"_dem"]}*(1+IF({L}{pv_rows["retard"]}=1,0,{g}*{L}{pv_rows["mult"]}))'
    pv_rows[key + '_dem'] = rw
    pv_line(rw, key + '_dem', f'Demande ({gmv_lo}) — croissance section 7 × multiplicateur acquisition',
            lambda L, Lp, i: f'={PR(dem_first_param)}*{STRESS_M1}', dem_next, CNT1); rw += 1
    pv_line(rw, key + '_real', f'{gmv_up} réalisées = MIN(demande, capacité)',
            lambda L, Lp, i: f'=MIN({L}{rw-1},{L}{pv_rows[key+"_cap"]})', lambda L, Lp, i: f'=MIN({L}{rw-1},{L}{pv_rows[key+"_cap"]})', CNT1, bold=True); rw += 1
    pv_line(rw, key + '_unmet', "Demande non servie (pénurie d'offre)",
            lambda L, Lp, i: f'={L}{rw-2}-{L}{rw-1}', lambda L, Lp, i: f'={L}{rw-2}-{L}{rw-1}', CNT1); rw += 1
    pv_line(rw, key + '_util', "Taux d'utilisation de la capacité brute (réalisées / (actifs × max))",
            lambda L, Lp, i: f'=IF({L}{pv_rows[key+"_stock"]}*{PR(maxk)}=0,0,{L}{rw-2}/({L}{pv_rows[key+"_stock"]}*{PR(maxk)}))',
            lambda L, Lp, i: f'=IF({L}{pv_rows[key+"_stock"]}*{PR(maxk)}=0,0,{L}{rw-2}/({L}{pv_rows[key+"_stock"]}*{PR(maxk)}))', PCT); rw += 1
    pv_line(rw, key + '_price', 'Panier / loyer moyen (indexé par an)',
            lambda L, Lp, i: f'={PR(price_param)}*{infl(infl_name, L)}', lambda L, Lp, i: f'={PR(price_param)}*{infl(infl_name, L)}'); rw += 1
    pv_line(rw, key + '_gmv', f'GMV {gmv_lo}', lambda L, Lp, i: f'={L}{pv_rows[key+"_real"]}*{L}{rw-1}', lambda L, Lp, i: f'={L}{pv_rows[key+"_real"]}*{L}{rw-1}', bold=True); rw += 1
    pv_line(rw, key + '_ca', ca_label, lambda L, Lp, i: take_expr(L, rw - 1), lambda L, Lp, i: take_expr(L, rw - 1), bold=True); rw += 1
    return rw + 1

# la ligne du flag collab 2 est fixée en fin de feuille ; on réserve son numéro maintenant
pv_rows['flag'] = 70

rw = supply_block(rw, "OFFRE & DEMANDE — PRESTATIONS AMPIO (prestataires, capacité, GMV, commissions, abonnements Ampio+)", 'p',
                  'prest_m1', 'prest_onb', 'prest_onb2', 'prest_churn', 'jobs_max', 'presta_m1', 'g_presta', 'cap_presta', 'panier', 'infl_var',
                  lambda L, gmvrow: f'=MAX(0,{L}{gmvrow}*({PR("take_presta")}-IF({MO(L)}<={PR("launch_dur")},{PR("launch_tr")},0)))',
                  ('prestations', 'Prestations'), 'CA commissions prestations (take rate mix − subvention de lancement M1–M6)',
                  ('Nouveaux prestataires onboardés', 'Prestataires perdus', 'Prestataires actifs'), with_abo=True)

# vidéo
section(PV, rw, 'MISE EN LIGNE VIDÉO — Ampianaro B2B (croissance ralentie A2/A3 × multiplicateur acquisition)', NCOL); rw += 1
pv_rows['v_stock'] = rw + 2
pv_line(rw, 'v_new', 'Nouveaux clients vidéo (croissance nette + remplacement du churn)',
        lambda L, Lp, i: f'={L}{rw+2}', lambda L, Lp, i: f'={L}{rw+2}-{Lp}{rw+2}*(1-{PR("churn_video")})', CNT1); rw += 1
pv_line(rw, 'v_lost', 'Clients perdus (churn)', lambda L, Lp, i: '=0', lambda L, Lp, i: f'={Lp}{rw+1}*{PR("churn_video")}', CNT1); rw += 1
pv_line(rw, 'v_stock', 'Clients vidéo actifs', lambda L, Lp, i: f'={PR("video_m1")}*{STRESS_M1}',
        lambda L, Lp, i: f'={Lp}{rw}*(1+IF({L}{pv_rows["retard"]}=1,0,{GROWTH("g_video", L)}*{L}{pv_rows["mult"]}))', CNT1, bold=True); rw += 1
pv_line(rw, 'v_price', 'Prix du service vidéo indexé', lambda L, Lp, i: f'={PR("prix_video")}*{infl("index_prix", L)}',
        lambda L, Lp, i: f'={PR("prix_video")}*{infl("index_prix", L)}'); rw += 1
pv_line(rw, 'v_ca', 'CA vidéo (MRR Ampianaro B2B)', lambda L, Lp, i: f'={L}{rw-2}*{L}{rw-1}', lambda L, Lp, i: f'={L}{rw-2}*{L}{rw-1}', bold=True); rw += 2

rw = supply_block(rw, 'OFFRE & DEMANDE — LOCATIONS PRO Ampindramo B2B2C (loueurs, capacité, paliers 30/20/15 %)', 'lp',
                  'loueurs_m1', 'loueurs_onb', 'loueurs_onb2', 'loueurs_churn', 'loc_max', 'locpro_m1', 'g_locpro', 'cap_locpro', 'loyer_pro', 'infl_var',
                  lambda L, gmvrow: f'={L}{gmvrow}*{PR("take_locpro")}', ('locations pro', 'Locations pro'), 'CA commissions locations pro',
                  ('Nouveaux loueurs pro onboardés', 'Loueurs pro perdus', 'Loueurs pro actifs'))
rw = supply_block(rw, 'OFFRE & DEMANDE — LOCATIONS PARTICULIERS Ampindramo C2C (annonces, capacité, 13 %)', 'c2',
                  'annonces_m1', 'annonces_onb', None, 'annonces_churn', 'ann_max', 'c2c_m1', 'g_c2c', None, 'loyer_c2c', 'infl_var',
                  lambda L, gmvrow: f'={L}{gmvrow}*{PR("take_c2c")}', ('locations C2C', 'Locations C2C'), 'CA commissions locations particuliers',
                  ('Nouvelles annonces C2C publiées', 'Annonces C2C retirées', 'Annonces C2C actives'))

section(PV, rw, 'TOTAUX', NCOL); rw += 1
def tot(key, label, expr, fmt=MONEY, bold=False):
    global rw
    pv_line(rw, key, label, lambda L, Lp, i: expr(L), lambda L, Lp, i: expr(L), fmt, bold); rw += 1
tot('gmv', 'GMV total (prestations + locations)', lambda L: f'={L}{pv_rows["p_gmv"]}+{L}{pv_rows["lp_gmv"]}+{L}{pv_rows["c2_gmv"]}', bold=True)
tot('mrr', 'MRR (abonnements Ampio+ + vidéo)', lambda L: f'={L}{pv_rows["ca_abo"]}+{L}{pv_rows["v_ca"]}', bold=True)
tot('ca_tx', 'CA transactionnel (commissions)', lambda L: f'={L}{pv_rows["p_ca"]}+{L}{pv_rows["lp_ca"]}+{L}{pv_rows["c2_ca"]}', bold=True)
tot('ca', "CHIFFRE D'AFFAIRES TOTAL (IFRS 15 — agent : commissions + abonnements)", lambda L: f'={L}{pv_rows["mrr"]}+{L}{pv_rows["ca_tx"]}', bold=True)
tot('tx', 'Transactions réalisées (prestations + locations pro + C2C)', lambda L: f'={L}{pv_rows["p_real"]}+{L}{pv_rows["lp_real"]}+{L}{pv_rows["c2_real"]}', CNT1)
tot('new_dem', 'Nouveaux clients DEMANDE estimés (part transactions nouveaux ÷ transactions par client)', lambda L: f'={PR("new_share")}*{L}{pv_rows["tx"]}/{PR("tx_per_new")}', CNT1)
tot('new_verif', 'Nouveaux clients à vérifier (prestataires + loueurs + vidéo + demande)', lambda L: f'={L}{pv_rows["p_new"]}+{L}{pv_rows["lp_new"]}+{L}{pv_rows["v_new"]}+{L}{pv_rows["new_dem"]}', CNT1)
tot('unmet_gmv', "GMV perdu faute d'offre (demande non servie × panier)", lambda L: f'={L}{pv_rows["p_unmet"]}*{L}{pv_rows["p_price"]}+{L}{pv_rows["lp_unmet"]}*{L}{pv_rows["lp_price"]}+{L}{pv_rows["c2_unmet"]}*{L}{pv_rows["c2_price"]}')
rw += 1
section(PV, rw, 'EMBAUCHE DU COLLABORATEUR 2', NCOL); rw += 1
assert rw == pv_rows['flag'], (rw, pv_rows['flag'])
FLAGROW = rw
def flag_next(L, Lp, i):
    T = f'Tresorerie!$C$25:$AM$25'
    PLR = "'P&L'!$B$18:$AK$18"
    m = MO(L)
    return (f"=IF({m}<{PR('c2_min_month')},0,IF({Lp}{FLAGROW}=1,1,IF(AND(N({m}>={PR('c2_consec')}+1),"
            f"INDEX({T},{m})>{PR('c2_plancher')},INDEX({T},{m}-1)>{PR('c2_plancher')},INDEX({T},{m}-2)>{PR('c2_plancher')},"
            f"IF({PR('c2_ebitda')}=1,INDEX({PLR},{m}-1)>0,TRUE),IF({PR('c2_ebitda')}=1,INDEX({PLR},{m}-2)>0,TRUE),IF({PR('c2_ebitda')}=1,INDEX({PLR},{m}-3)>0,TRUE)),1,0)))")
pv_line(rw, 'flag', 'Collaborateur 2 actif (0/1) — déclencheur trésorerie + EBITDA, mois minimum section 14', lambda L, Lp, i: '=0', flag_next, INT, bold=True)
PV['B' + str(FLAGROW)].font = Font(name='Calibri', size=10, bold=True, color='0000FF')
PV.column_dimensions['A'].width = 78
for i in range(2, NCOL + 1): PV.column_dimensions[CL(i)].width = 11
PV.freeze_panes = 'B4'; PV.sheet_view.showGridLines = False
def PVR(key, L='B'): return f"'Plan ventes'!{L}{pv_rows[key]}"

# ---------------------------------------------------------------- P&L : lignes 5, 6, 7, 12, 16 + libellés
PL = wb['P&L']
for i in range(1, 37):
    L = CL(1 + i)
    PL[f'{L}5'] = f"={PVR('ca', L)}-IF({PR('st_on','B')}=1,{PR('st_takerate','B')}*{PVR('gmv', L)},0)"
    PL[f'{L}6'] = f"={PVR('mrr', L)}"
    PL[f'{L}7'] = (f"=-(({PVR('ca_abo',L)}+{PVR('p_gmv',L)}+{PVR('lp_gmv',L)}+{PVR('c2_gmv',L)})*{PR('mm')}"
                   f"+({PVR('p_real',L)}+{PVR('c2_real',L)})*{PR('sms')}*{infl('infl_var', L)}"
                   f"+({PVR('tx',L)})*{PR('support')}+{PVR('new_verif',L)}*{PR('verif')}+{PVR('gmv',L)}*{PR('litiges')})")
    PL[f'{L}12'] = (f"=-((({PR('sal_fond')}+{PR('sal_c1')})*(1+{PR('charges_taux')}))*{infl('infl_opex', L)}"
                    f"+IF({PVR('flag',L)}=1,{PR('sal_c2')}*(1+{PR('charges_taux')})*{infl('infl_opex', L)},0))")
    PL[f'{L}16'] = f"=-({PR('accomp')}*{infl('infl_opex', L)}+{PVR('pub', L)})"
    for rr_ in (5, 6, 7, 12, 16):
        PL[f'{L}{rr_}'].font = fcalc(PL[f'{L}{rr_}'].value, bold=PL[f'{L}{rr_}'].font.bold); WRITTEN.add(('P&L', f'{L}{rr_}'))
PL['A7'] = 'Coûts variables (Mobile Money + SMS + support + vérification + litiges)'
PL['A12'] = 'Salaires chargés — fondateur + collab 1 (+ collab 2 si actif), charges patronales incluses'
PL['A16'] = 'Marketing (accompagnement fixe + budget publicité par année)'
PL['A24'] = 'Impôt (IS synthétique 5 % du CA ou IR 20 %) — comptabilisé au mois 12 de chaque exercice'

# ---------------------------------------------------------------- Contribution : lignes 6-10, 15, note
CO = wb['Contribution']
for i in range(1, 37):
    L = CL(1 + i)
    CO[f'{L}6'] = f"=-({PVR('ca_abo',L)}+{PVR('p_gmv',L)}+{PVR('lp_gmv',L)}+{PVR('c2_gmv',L)})*{PR('mm')}"
    CO[f'{L}7'] = f"=-({PVR('p_real',L)}+{PVR('c2_real',L)})*{PR('sms')}*{infl('infl_var', L)}"
    CO[f'{L}8'] = f"=-{PVR('tx',L)}*{PR('support')}"
    CO[f'{L}9'] = f"=-{PVR('new_verif',L)}*{PR('verif')}"
    CO[f'{L}10'] = f"=-{PVR('gmv',L)}*{PR('litiges')}"
    CO[f'{L}15'] = f"=-{PVR('pub',L)}"
    for rr_ in (6, 7, 8, 9, 10, 15):
        CO[f'{L}{rr_}'].font = fcalc(CO[f'{L}{rr_}'].value); WRITTEN.add(('Contribution', f'{L}{rr_}'))
CO['A9'] = 'dont vérification nouveaux clients (prestataires, loueurs, vidéo, demande)'
CO['A15'] = 'Publicité (budget saisi par année, Parametres section 9)'
CO['A18'] = 'OPEX fixes (loyer, salaires chargés, hébergement, storage, modération, accompagnement)'
CO['A22'] = ("Lecture : les coûts variables suivent le volume (transactions réalisées, nouveaux clients, GMV). Les OPEX fixes ne bougent qu'avec "
             "l'inflation, le scénario, le stress test et l'embauche du collab 2. Taux unitaires : Parametres section 12 ; budget pub : section 9.")

# ---------------------------------------------------------------- Base graphique
BG = wb['Base graphique']
for i in range(1, 37):
    L = CL(1 + i)
    BG[f'{L}6'] = f"={PVR('mrr', L)}"; BG[f'{L}7'] = f"={PVR('gmv', L)}"
    WRITTEN.add(('Base graphique', f'{L}6')); WRITTEN.add(('Base graphique', f'{L}7'))
do_remap_keep()

# ---------------------------------------------------------------- FINANCEMENT (regénéré avec bloc Ask)
F = wb['Financement']
FMAP = {4: 4, 11: 11, 14: 14, 15: 15, 20: 23, 21: 24, 22: 25, 28: 31, 29: 32, 30: 33, 31: 34}
fref = re.compile(r"Financement!(\$?)([A-Z]{1,2})(\$?)(\d{1,3})")
for sn in ('Bilan', 'Tresorerie', 'P&L'):
    for row in wb[sn].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and 'Financement!' in c.value:
                c.value = fref.sub(lambda m: f"Financement!{m.group(1)}{m.group(2)}{m.group(3)}{FMAP[int(m.group(4))]}", c.value)
clear_sheet(F, rows=50, cols=38)
LASTF = 38
def fline(r_, lab, formula, fmt=MONEY, bold=False, note=None, total=False):
    a = F.cell(row=r_, column=1, value=lab); a.font = Font(name='Calibri', size=10, bold=bold)
    b = F.cell(row=r_, column=2, value=formula); b.font = fcalc(formula, bold); b.number_format = fmt
    if total: a.border = B_TOP; b.border = B_TOP
    if note:
        n = F.cell(row=r_, column=3, value=note); n.font = Font(name='Calibri', size=9, italic=True, color=C_GRAY); n.alignment = Alignment(vertical='top')
title(F, "FINANCEMENT — besoin initial, plan de financement (Ask), emprunt et validation sur 3 ans", LASTF)
section(F, 3, 'BESOIN DE FINANCEMENT INITIAL', LASTF)
fline(4, 'CAPEX (investissements, choc stress inclus)', f"={PR('total_capex')}*(1+IF({PR('st_on','B')}=1,{PR('st_capex','B')},0))")
fline(5, 'Charges fixes mensuelles année 1 (OPEX chargés + budget publicité A1)', f"={PR('total_opex')}+{PR('pub_a1')}", note='Parametres sections 13 et 9 — multiplicateur scénario et choc stress inclus.')
fline(6, 'Coûts variables du mois 1', "=-'P&L'!B7", note='Mobile Money, SMS, support, vérification, litiges du premier mois.')
fline(7, 'Nombre de mois de démarrage couverts', f"={PR('fin_mois_dem')}", fmt=INT)
fline(8, 'Besoin de démarrage (N mois × (charges fixes + coûts variables M1))', "=B7*(B5+B6)", bold=True)
fline(9, "Nombre de mois de charges fixes en sécurité", f"={PR('fin_mois_sec')}", fmt=INT)
fline(10, "Trésorerie de sécurité (N mois × charges fixes) — cible de la validation", "=B9*B5", bold=True)
fline(11, 'TOTAL BESOIN DE FINANCEMENT', "=B4+B8+B10", bold=True, total=True)
section(F, 13, "PLAN DE FINANCEMENT & ASK (demande de fonds)", LASTF)
fline(14, 'Apport en fonds propres (fondateurs)', f"={PR('apport')}")
fline(15, "Emprunt (mécanique inchangée : besoin non couvert par l'apport, porté en dette dans la trésorerie)", "=MAX(B11-B14,0)", bold=True)
fline(16, 'TOTAL RESSOURCES (= besoin)', "=B14+B15", bold=True, total=True)
fline(17, "Part du financement externe dans le total", "=IF(B16=0,0,B15/B16)", fmt=PCT)
fline(18, "ASK — financement externe recherché auprès d'investisseurs / bailleurs (besoin − apport)", "=B11-B14", bold=True,
      note="Montant à présenter aux financeurs. Ventilation indicative ci-dessous ; la trésorerie modélise l'ensemble en emprunt.")
fline(19, "   dont visé en equity / quasi-fonds propres (investisseurs) — indicatif", f"=B18*{PR('ask_inv')}", note='Part paramétrable (Parametres section 16).')
fline(20, "   dont visé en dette (banques / bailleurs) — indicatif", "=B18-B19")
section(F, 22, "EMPRUNT — CONDITIONS", LASTF)
fline(23, 'Montant emprunté', "=B15")
fline(24, 'Taux annuel', f"={PR('taux')}", fmt=PCT)
fline(25, 'Durée (mois)', f"={PR('duree')}", fmt=INT)
fline(26, 'Mensualité constante (annuités)', "=IF(B23=0,0,-PMT(B24/12,B25,B23))")
fline(27, "Coût total des intérêts sur la durée", "=B32", note="Intérêts → P&L ligne 22 et Trésorerie ligne 9 ; capital → Trésorerie ligne 20 ; dette restante → Bilan ligne 13.")
section(F, 29, "TABLEAU D'AMORTISSEMENT DE L'EMPRUNT (mensuel)", LASTF)
F.cell(row=30, column=1, value='Mois'); F.cell(row=30, column=2, value='TOTAL')
for i in range(1, 37): F.cell(row=30, column=2 + i, value=i).number_format = INT
for c in range(1, LASTF + 1):
    cl = F.cell(row=30, column=c); cl.font = F_HDR; cl.fill = FILL_H
for r_, lab in {31: 'Mensualité', 32: 'Intérêts', 33: 'Capital remboursé', 34: 'Capital restant dû'}.items():
    F.cell(row=r_, column=1, value=lab).font = Font(name='Calibri', size=10, bold=(r_ in (31, 34)))
for r_ in (31, 32, 33):
    b = F.cell(row=r_, column=2, value=f'=SUM(C{r_}:AL{r_})'); b.font = Font(name='Calibri', size=10, bold=True); b.number_format = MONEY
for i in range(1, 37):
    L = CL(2 + i)
    F[f'{L}31'] = f'=IF({L}$30>$B$25,0,-PMT($B$24/12,$B$25,$B$23))'
    F[f'{L}32'] = f'=IF({L}$30>$B$25,0,-IPMT($B$24/12,{L}$30,$B$25,$B$23))'
    F[f'{L}33'] = f'=IF({L}$30>$B$25,0,-PPMT($B$24/12,{L}$30,$B$25,$B$23))'
    F[f'{L}34'] = '=$B$23-C33' if i == 1 else f'={CL(1 + i)}34-{L}33'
    for r_ in (31, 32, 33, 34):
        F[f'{L}{r_}'].number_format = MONEY; F[f'{L}{r_}'].font = Font(name='Calibri', size=10, bold=(r_ == 34))
section(F, 36, "VALIDATION 3 ANS — la trésorerie reste-t-elle au-dessus du matelas de sécurité sur M0–M36 ?", LASTF)
fline(37, 'Pic de besoin cumulé de trésorerie AVANT financement (flux op. hors intérêts + CAPEX)', "=-MIN(B48:AL48)", bold=True,
      note="Point bas du cumul des flux (ligne 48) : montant à financer pour ne jamais être à zéro, hors coût de la dette.")
fline(38, "Financement minimal pour conserver le matelas de sécurité (pic + sécurité, hors service de la dette)", "=B37+B10")
fline(39, 'Trésorerie minimale sur 36 mois APRÈS financement et service de la dette', "=MIN(Tresorerie!D25:AM25)", bold=True)
fline(40, 'Couverture du matelas de sécurité (cible = ligne 10)',
      f'=IF(B39>=B10,"OK — trésorerie ≥ "&B9&" mois de charges fixes sur toute la période","INSUFFISANT — augmenter apport, mois financés ou durée")', fmt='General', bold=True)
fline(41, "Financement externe minimum indicatif pour atteindre la cible (itération suivante, hors intérêts additionnels)", "=B15+MAX(0,B10-B39)")
fline(42, 'Référence V2.3 — apport 36,7 M Ar calé à la main pour un plancher de 1 M', f"={PR('ref_v23')}")
fline(43, 'Écart référence 36,7 M vs financement minimal (positif = référence surdimensionnée)', "=B42-B38")
fline(44, "Capital restant dû à la fin du mois 36 (emprunt plus long que l'horizon du modèle)", "=AL34", bold=True, note="Nul si la durée ≤ 36 mois.")
fline(45, 'Trésorerie fin M36 nette de la dette restante', "=Tresorerie!AM25-B44", bold=True)
F.cell(row=48, column=1, value='Cumul flux opérationnels (hors intérêts) + CAPEX, avant financement — M0 (col. B) → M36 (col. AL)').font = Font(name='Calibri', size=9, italic=True, color=C_GRAY)
for i in range(0, 37):
    L, T = CL(2 + i), CL(3 + i)
    flux = f"(SUM(Tresorerie!{T}6:{T}8)+Tresorerie!{T}10+Tresorerie!{T}15)"
    F[f'{L}48'] = f'={flux}' if i == 0 else f'={CL(1 + i)}48+{flux}'
    F[f'{L}48'].number_format = MONEY; F[f'{L}48'].font = Font(name='Calibri', size=9, color='008000')
F.column_dimensions['A'].width = 84; F.column_dimensions['B'].width = 16
for i in range(3, LASTF + 1): F.column_dimensions[CL(i)].width = 11
F.freeze_panes = None; F.sheet_view.showGridLines = False

# ---------------------------------------------------------------- KPI (regénéré)
K = wb['KPI']
clear_sheet(K, rows=90, cols=37)
title(K, 'KPI MARKETPLACE — mensuels M1–M36 (MGA) et synthèse annuelle', NCOL)
hdr_row(K, 2, ['Année'], NCOL); hdr_row(K, 3, ['Mois'], NCOL)
for i in range(1, 37):
    L = CL(1 + i)
    K[f'{L}2'] = f"='Plan ventes'!{L}2"; K[f'{L}3'] = f"='Plan ventes'!{L}3"
    K[f'{L}2'].number_format = INT; K[f'{L}3'].number_format = INT
krow = {}
def kline(r_, key, label, expr, fmt=MONEY, bold=False):
    krow[key] = r_
    K.cell(row=r_, column=1, value=label).font = Font(name='Calibri', size=10, bold=bold)
    for i in range(1, 37):
        L, Lp = CL(1 + i), CL(i)
        f = expr(L, Lp, i)
        c = K.cell(row=r_, column=1 + i, value=f); c.number_format = fmt; c.font = fcalc(f, bold)
def KR(key, L): return f'{L}{krow[key]}'
rk = 5
section(K, rk, "CROISSANCE & VOLUME D'AFFAIRES", NCOL); rk += 1
kline(rk, 'gmv', "GMV (volume d'affaires brut)", lambda L, Lp, i: f"={PVR('gmv', L)}"); rk += 1
kline(rk, 'ca', "Chiffre d'affaires net (IFRS 15, après stress)", lambda L, Lp, i: f"='P&L'!{L}5"); rk += 1
kline(rk, 'tr', 'Take rate effectif (CA transactionnel / GMV)', lambda L, Lp, i: f"=IF({PVR('gmv', L)}=0,0,{PVR('ca_tx', L)}/{PVR('gmv', L)})", PCT); rk += 1
kline(rk, 'mrr', 'MRR (abonnements + vidéo)', lambda L, Lp, i: f"={PVR('mrr', L)}"); rk += 1
kline(rk, 'arr', 'ARR (MRR × 12)', lambda L, Lp, i: f"={KR('mrr', L)}*12"); rk += 1
kline(rk, 'g_mom', 'Croissance CA (MoM)', lambda L, Lp, i: '=0' if i == 1 else f"=IF({KR('ca', Lp)}=0,0,{KR('ca', L)}/{KR('ca', Lp)}-1)", PCT); rk += 1
kline(rk, 'pub', 'Budget publicité du mois', lambda L, Lp, i: f"={PVR('pub', L)}"); rk += 1
kline(rk, 'mult', "Multiplicateur de croissance (acquisition)", lambda L, Lp, i: f"={PVR('mult', L)}", RATIO); rk += 2
section(K, rk, "OFFRE & LIQUIDITÉ (capacité contraignante)", NCOL); rk += 1
kline(rk, 'prest', 'Prestataires Ampio actifs', lambda L, Lp, i: f"={PVR('p_stock', L)}", CNT1); rk += 1
kline(rk, 'util_p', "Utilisation de la capacité Ampio (alerte si ≥ fill rate max)", lambda L, Lp, i: f"={PVR('p_util', L)}", PCT); rk += 1
kline(rk, 'unmet', "Prestations non servies (pénurie d'offre)", lambda L, Lp, i: f"={PVR('p_unmet', L)}", CNT1); rk += 1
kline(rk, 'unmet_gmv', "GMV perdu faute d'offre (tous segments)", lambda L, Lp, i: f"={PVR('unmet_gmv', L)}"); rk += 1
kline(rk, 'loueurs', 'Loueurs pro actifs', lambda L, Lp, i: f"={PVR('lp_stock', L)}", CNT1); rk += 1
kline(rk, 'util_lp', 'Utilisation capacité locations pro', lambda L, Lp, i: f"={PVR('lp_util', L)}", PCT); rk += 1
kline(rk, 'annonces', 'Annonces C2C actives', lambda L, Lp, i: f"={PVR('c2_stock', L)}", CNT1); rk += 1
kline(rk, 'util_c2', 'Utilisation capacité C2C', lambda L, Lp, i: f"={PVR('c2_util', L)}", PCT); rk += 1
kline(rk, 'flag', 'Collaborateur 2 actif (0/1)', lambda L, Lp, i: f"={PVR('flag', L)}", INT); rk += 2
section(K, rk, "UNIT ECONOMICS PAR FACE (offre / demande) — LTV bornée à 36 mois, commissions partagées 50/50 entre les faces", NCOL); rk += 1
kline(rk, 'new_off', 'Nouveaux OFFRE (prestataires + loueurs pro onboardés)', lambda L, Lp, i: f"={PVR('p_new', L)}+{PVR('lp_new', L)}", CNT1); rk += 1
kline(rk, 'new_dem', 'Nouveaux DEMANDE (clients vidéo + clients transactionnels estimés)', lambda L, Lp, i: f"={PVR('v_new', L)}+{PVR('new_dem', L)}", CNT1); rk += 1
kline(rk, 'spend', "Dépenses d'acquisition (pub + part accompagnement + temps fondateur chargé)",
      lambda L, Lp, i: f"={PVR('pub', L)}+{PR('accomp')}*{infl('infl_opex', L)}*{PR('acq_share')}+{PR('founder_time')}*{PR('sal_fond')}*(1+{PR('charges_taux')})*{infl('infl_opex', L)}"); rk += 1
kline(rk, 'spend_off', "Dépenses d'acquisition OFFRE", lambda L, Lp, i: f"={PR('split_off')}*{KR('spend', L)}"); rk += 1
kline(rk, 'spend_dem', "Dépenses d'acquisition DEMANDE", lambda L, Lp, i: f"={PR('split_dem')}*{KR('spend', L)}"); rk += 1
kline(rk, 'cac_off', 'CAC offre', lambda L, Lp, i: f"=IF({KR('new_off', L)}=0,0,{KR('spend_off', L)}/{KR('new_off', L)})"); rk += 1
kline(rk, 'cac_dem', 'CAC demande', lambda L, Lp, i: f"=IF({KR('new_dem', L)}=0,0,{KR('spend_dem', L)}/{KR('new_dem', L)})"); rk += 1
kline(rk, 'arpu_off', 'Revenu mensuel moyen par prestataire / loueur (50 % commissions + abonnements Ampio+)',
      lambda L, Lp, i: f"=IF({PVR('p_stock', L)}+{PVR('lp_stock', L)}=0,0,(0.5*({PVR('p_ca', L)}+{PVR('lp_ca', L)})+{PVR('ca_abo', L)})/({PVR('p_stock', L)}+{PVR('lp_stock', L)}))"); rk += 1
kline(rk, 'life_off', "Durée de vie offre (mois) = MIN(36 ; 1 / churn prestataires)", lambda L, Lp, i: f"=MIN(36,IF({PR('prest_churn')}=0,36,1/{PR('prest_churn')}))", CNT1); rk += 1
kline(rk, 'ltv_off', 'LTV offre (revenu mensuel × durée de vie bornée)', lambda L, Lp, i: f"={KR('arpu_off', L)}*{KR('life_off', L)}"); rk += 1
kline(rk, 'ltv_dem', 'LTV demande (vidéo : prix × MIN(36 ; 1/churn) ; transactionnel : 50 % commission par transaction × tx/client × durée de vie)',
      lambda L, Lp, i: (f"=IF({KR('new_dem', L)}=0,0,({PVR('v_new', L)}*{PVR('v_price', L)}*MIN(36,IF({PR('churn_video')}=0,36,1/{PR('churn_video')}))"
                        f"+{PVR('new_dem', L)}*0.5*IF({PVR('tx', L)}=0,0,{PVR('ca_tx', L)}/{PVR('tx', L)})*{PR('tx_per_new')}*{PR('life_tx')})/{KR('new_dem', L)})")); rk += 1
kline(rk, 'ltvcac_off', 'LTV / CAC offre (cible indicative > 3)', lambda L, Lp, i: f"=IF({KR('cac_off', L)}=0,0,{KR('ltv_off', L)}/{KR('cac_off', L)})", RATIO); rk += 1
kline(rk, 'ltvcac_dem', 'LTV / CAC demande (cible indicative > 3)', lambda L, Lp, i: f"=IF({KR('cac_dem', L)}=0,0,{KR('ltv_dem', L)}/{KR('cac_dem', L)})", RATIO); rk += 1
kline(rk, 'payback', 'Payback CAC offre (mois) = CAC / revenu mensuel par prestataire', lambda L, Lp, i: f"=IF({KR('arpu_off', L)}=0,0,{KR('cac_off', L)}/{KR('arpu_off', L)})", CNT1); rk += 2
section(K, rk, 'TRÉSORERIE & LIQUIDITÉ FINANCIÈRE', NCOL); rk += 1
kline(rk, 'cash', 'Trésorerie de clôture', lambda L, Lp, i: f"=Tresorerie!{CL(3 + i)}25"); rk += 1
kline(rk, 'burn', 'Burn net mensuel', lambda L, Lp, i: f"=Tresorerie!{CL(3 + i)}27"); rk += 1
kline(rk, 'runway', 'Runway (mois)', lambda L, Lp, i: f"=Tresorerie!{CL(3 + i)}28", CNT1); rk += 1
kline(rk, 'pc', 'Passif courant estimé (capital emprunt ≤ 12 mois + impôt de l\'exercice)',
      lambda L, Lp, i: f"=MIN(Financement!{CL(2 + i)}34,12*Financement!$C$31)+CHOOSE('P&L'!{L}2,'P&L'!$B$34,'P&L'!$C$34,'P&L'!$D$34)"); rk += 1
kline(rk, 'liq', 'Taux de liquidité générale (trésorerie / passif courant)', lambda L, Lp, i: f"=IF({KR('pc', L)}=0,\"∞\",{KR('cash', L)}/{KR('pc', L)})", RATIO); rk += 2
section(K, rk, 'SYNTHÈSE ANNUELLE', NCOL); rk += 1
hdr_row(K, rk, [None, 'Année 1', 'Année 2', 'Année 3'], 4); rk += 1
arow = {}
def aline(key, label, exprs, fmt=MONEY, bold=False):
    global rk
    arow[key] = rk
    K.cell(row=rk, column=1, value=label).font = Font(name='Calibri', size=10, bold=bold)
    for j, col in enumerate('BCD'):
        f = exprs(j + 1, col)
        c = K.cell(row=rk, column=2 + j, value=f); c.number_format = fmt; c.font = fcalc(f, bold)
    rk += 1
def SUMY(key, y): return f"SUMIF($B$2:$AK$2,{y},$B${krow[key]}:$AK${krow[key]})"
def ENDY(key, y): return f"INDEX($B${krow[key]}:$AK${krow[key]},1,12*{y})"
aline('gmv', 'GMV annuel', lambda y, c: f"={SUMY('gmv', y)}")
aline('ca', 'CA annuel', lambda y, c: f"='P&L'!{c}30")
aline('tr', 'Take rate moyen (CA transactionnel / GMV)', lambda y, c: f"=IF({SUMY('gmv', y)}=0,0,SUMIF($B$2:$AK$2,{y},'Plan ventes'!$B${pv_rows['ca_tx']}:$AK${pv_rows['ca_tx']})/{SUMY('gmv', y)})", PCT)
aline('mrr', "MRR fin d'année", lambda y, c: f"={ENDY('mrr', y)}")
aline('ebitda', 'EBITDA annuel', lambda y, c: f"='P&L'!{c}31")
aline('rn', 'Résultat net annuel', lambda y, c: f"='P&L'!{c}35")
aline('m_ebitda', 'Marge EBITDA', lambda y, c: f"=IF({c}{arow['ca']}=0,0,{c}{arow['ebitda']}/{c}{arow['ca']})", PCT)
aline('g_yoy', 'Croissance CA (YoY)', lambda y, c: '=0' if y == 1 else f"=IF({CL(y)}{arow['ca']}=0,0,{c}{arow['ca']}/{CL(y)}{arow['ca']}-1)", PCT)
aline('prest', "Prestataires actifs fin d'année", lambda y, c: f"={ENDY('prest', y)}", CNT1)
aline('util', "Utilisation capacité Ampio fin d'année", lambda y, c: f"={ENDY('util_p', y)}", PCT)
aline('unmet_gmv', "GMV perdu faute d'offre (annuel)", lambda y, c: f"={SUMY('unmet_gmv', y)}")
aline('pub', 'Budget publicité annuel', lambda y, c: f"={SUMY('pub', y)}")
aline('cac_off', 'CAC offre moyen annuel', lambda y, c: f"=IF({SUMY('new_off', y)}=0,0,{SUMY('spend_off', y)}/{SUMY('new_off', y)})")
aline('cac_dem', 'CAC demande moyen annuel', lambda y, c: f"=IF({SUMY('new_dem', y)}=0,0,{SUMY('spend_dem', y)}/{SUMY('new_dem', y)})")
aline('ltvcac_off', "LTV / CAC offre fin d'année", lambda y, c: f"={ENDY('ltvcac_off', y)}", RATIO)
aline('ltvcac_dem', "LTV / CAC demande fin d'année", lambda y, c: f"={ENDY('ltvcac_dem', y)}", RATIO)
aline('r40', 'Rule of 40 (croissance CA YoY + marge EBITDA)', lambda y, c: f"={c}{arow['g_yoy']}+{c}{arow['m_ebitda']}", PCT)
K.column_dimensions['A'].width = 80
for i in range(2, NCOL + 1): K.column_dimensions[CL(i)].width = 11
K.freeze_panes = 'B4'; K.sheet_view.showGridLines = False

# ---------------------------------------------------------------- DASHBOARD (colonnes A:B ; graphiques conservés en F:X)
D = wb['Dashboard']
for r_ in range(2, 60):
    for c in (1, 2):
        cl = D.cell(row=r_, column=c); cl.value = None; cl.font = Font(name='Calibri', size=10); cl.fill = PatternFill(); cl.number_format = 'General'; cl.border = Border()
AR = '#,##0" Ar"'
def dsec(r_, text):
    D.cell(row=r_, column=1, value=text)
    for c in (1, 2):
        cl = D.cell(row=r_, column=c); cl.font = F_SEC; cl.fill = FILL_S
def dline(r_, lab, f, fmt=AR, bold=False):
    D.cell(row=r_, column=1, value=lab).font = Font(name='Calibri', size=10, bold=bold)
    c = D.cell(row=r_, column=2, value=f); c.font = fcalc(f, bold); c.number_format = fmt
D['A2'] = 'Scénario actif :'; D['A2'].font = Font(name='Calibri', size=10, bold=True)
D['B2'] = '=Parametres!B3'; D['B2'].font = Font(name='Calibri', size=11, bold=True, color='008000')
KA = lambda key, c: f"KPI!{c}{arow[key]}"
rd = 4
dsec(rd, 'ACTIVITÉ'); rd += 1
for y, c in ((1, 'B'), (2, 'C'), (3, 'D')):
    dline(rd, f"Chiffre d'affaires Année {y}", f"='P&L'!{c}30"); rd += 1
dline(rd, 'Croissance CA A2 / A1', f"={KA('g_yoy','C')}", PCT); rd += 1
dline(rd, 'Croissance CA A3 / A2', f"={KA('g_yoy','D')}", PCT); rd += 1
dline(rd, 'GMV Année 3', f"={KA('gmv','D')}"); rd += 1
dline(rd, 'Take rate moyen Année 3', f"={KA('tr','D')}", PCT); rd += 1
dline(rd, 'MRR fin Année 3', f"={KA('mrr','D')}"); rd += 1
dline(rd, 'ARR fin Année 3', f"={KA('mrr','D')}*12"); rd += 2
dsec(rd, 'RENTABILITÉ'); rd += 1
dline(rd, 'Marge de contribution Année 3', "=IFERROR(SUM(Contribution!Z12:AK12)/SUM(Contribution!Z5:AK5),0)", PCT); rd += 1
dline(rd, 'EBITDA Année 3', "='P&L'!D31"); rd += 1
for y, c in ((1, 'B'), (2, 'C'), (3, 'D')):
    dline(rd, f'Résultat net Année {y}', f"='P&L'!{c}35"); rd += 1
dline(rd, 'Seuil de rentabilité mensuel (CA, moyenne 36 mois)', "=-AVERAGE('P&L'!B17:AK17)/AVERAGE('P&L'!B9:AK9)"); rd += 1
dline(rd, 'Premier mois EBITDA > 0', "=IFERROR(MATCH(TRUE(),INDEX('P&L'!B18:AK18>0,0),0),\"> 36 mois\")", INT); rd += 1
dline(rd, 'Premier mois de résultat net cumulé positif', "=IFERROR(MATCH(TRUE(),INDEX('Base graphique'!B8:AK8>0,0),0),\"> 36 mois\")", INT); rd += 2
dsec(rd, 'UNIT ECONOMICS & OFFRE'); rd += 1
dline(rd, 'CAC offre Année 3', f"={KA('cac_off','D')}"); rd += 1
dline(rd, 'CAC demande Année 3', f"={KA('cac_dem','D')}"); rd += 1
dline(rd, 'LTV / CAC offre fin A3 (cible > 3, LTV bornée 36 mois)', f"={KA('ltvcac_off','D')}", RATIO); rd += 1
dline(rd, 'LTV / CAC demande fin A3 (cible > 3)', f"={KA('ltvcac_dem','D')}", RATIO); rd += 1
dline(rd, 'Prestataires Ampio actifs fin A3', f"={KA('prest','D')}", CNT1); rd += 1
dline(rd, 'Utilisation capacité Ampio fin A3 (alerte si ≥ fill rate max)', f"={KA('util','D')}", PCT); rd += 1
dline(rd, "GMV perdu faute d'offre sur 3 ans", f"=SUM(KPI!B{arow['unmet_gmv']}:D{arow['unmet_gmv']})"); rd += 1
dline(rd, "Mois d'activation collaborateur 2 (ou > 36)", f"=IFERROR(MATCH(1,'Plan ventes'!B{FLAGROW}:AK{FLAGROW},0),\"> 36 mois\")", INT); rd += 2
dsec(rd, 'FINANCEMENT & TRÉSORERIE'); rd += 1
dline(rd, 'Besoin de financement initial (CAPEX + démarrage + sécurité)', "=Financement!B11"); rd += 1
dline(rd, 'Apport en fonds propres', "=Financement!B14"); rd += 1
dline(rd, 'ASK — financement externe recherché', "=Financement!B18", bold=True); rd += 1
dline(rd, '   dont equity / quasi-fonds propres (indicatif)', "=Financement!B19"); rd += 1
dline(rd, '   dont dette (indicatif) — modélisé en emprunt', "=Financement!B20"); rd += 1
dline(rd, 'Mensualité de l\'emprunt', "=Financement!B26"); rd += 1
dline(rd, 'Trésorerie fin Année 3', "=Tresorerie!AM25"); rd += 1
dline(rd, 'Trésorerie minimale sur 36 mois', "=Financement!B39"); rd += 1
dline(rd, 'Cible : matelas de sécurité', "=Financement!B10"); rd += 1
dline(rd, 'Couverture du matelas de sécurité', "=Financement!B40", 'General', bold=True); rd += 1
dline(rd, 'Capital restant dû fin M36', "=Financement!B44"); rd += 1
dline(rd, 'Taux de liquidité fin A3 (cible > 1)', f"=INDEX(KPI!B{krow['liq']}:AK{krow['liq']},1,36)", RATIO); rd += 2
dsec(rd, 'CONTRÔLES'); rd += 1
dline(rd, 'Rule of 40 Année 3 (indicatif, base petite)', f"={KA('r40','D')}", PCT); rd += 1
dline(rd, 'Contrôle bilan max |écart| (Ar)', "=MAX(ABS(Bilan!C18),ABS(Bilan!AM18))", '#,##0'); rd += 1
dline(rd, 'Contrôle contribution vs P&L max |écart| (Ar)', "=MAX(ABS(MIN(Contribution!B20:AK20)),ABS(MAX(Contribution!B20:AK20)))", '#,##0'); rd += 1
dline(rd, 'Contrôle mix take rate prestations (doit = 100 %)', f"={PR('ctrl_mix')}", PCT); rd += 1
D.column_dimensions['A'].width = 58; D.column_dimensions['B'].width = 20
D.freeze_panes = 'A3'

# ---------------------------------------------------------------- Lisez-moi
LM = wb['Lisez-moi']
LM['A8'] = "1. Abonnement prestataire Ampio+ : 24 000 Ar/mois (indexé), souscrit par une part paramétrable des prestataires actifs (10 % par défaut)."
LM['A15'] = ("6. Coûts variables : commission Mobile Money 3 % du montant de chaque transaction (abonnement Ampio+, GMV prestations, GMV locations). Exclu : service vidéo. "
             "SMS (2 × 35,71 Ar) : prestations Ampio et locations C2C. Support 1 000 Ar/transaction, vérification 3 000 Ar/nouveau client, litiges 1 % du GMV.")
LM['A33'] = '2. Cellules à texte BLEU sur fond JAUNE : hypothèses modifiables ; texte NOIR : calculs ; texte VERT : références vers un autre onglet.'
LM['A36'] = 'CHANGEMENTS DE CETTE VERSION (refonte offre / acquisition / coûts)'
LM['A37'] = ("— Offre modélisée par onboarding et churn (prestataires, loueurs pro, annonces C2C) ; la capacité (actifs × max × fill rate max) PLAFONNE les volumes réalisés. "
             "Les anciens KPI de liquidité tautologiques (offre proportionnelle au volume, fill rate saisi) sont supprimés.")
LM['A38'] = "— Ampio+ fusionné avec l'offre : abonnés = pénétration (10 %) × prestataires actifs. Plus de stock d'abonnés ni de churn Ampio+ autonomes."
LM['A39'] = ("— Acquisition → croissance : budget publicité saisi par année (400 k / 500 k / 600 k Ar/mois) ; la croissance mensuelle = taux de base × "
             "(budget du mois précédent / budget de référence)^élasticité (0,5), borné 0,5–1,5. Le cold-start (plancher pub, règle 6 % du CA) est retiré ; "
             "la subvention d'offre (take rate −2 points M1–M6) est conservée et son bug corrigé (saisie 2 → 0,02).")
LM['A40'] = "— Ralentissement après A1 : vidéo 5 % → 4 % → 3 % ; prestations 2 % → 2,5 % → 2 % (l'accélération à 3 % est retirée)."
LM['A41'] = "— Indexation annuelle des prix (abonnement, vidéo) : 8 %, alignée sur l'inflation des coûts (nouveau paramètre, section 6)."
LM['A42'] = "— Charges patronales 18 % (CNaPS 13 % + organisation sanitaire 5 %) sur tous les salaires, incluses dans les OPEX et le besoin de financement."
LM['A43'] = "— Unit economics : LTV bornée à 36 mois ; commissions partagées 50/50 entre offre et demande ; CAC offre sur les prestataires et loueurs réellement onboardés."
LM['A44'] = "— Financement : bloc ASK (financement externe recherché = besoin − apport, ventilation indicative equity / dette) ; sécurité rétablie à 3 mois de charges fixes ; libellé de couverture dynamique."
LM['A45'] = "— Corrections techniques : le choc OPEX du stress test pointait vers des cellules déplacées (désactivé silencieusement) ; idem pour le choc CAPEX dans Financement. Toutes les formules sont régénérées à partir de paramètres nommés."
LM['A46'] = "— Plancher IR 1 M + 1 % CA ; impôt comptabilisé et payé au mois 12 de chaque exercice (inchangé)."
LM['A62'] = "4. Acquisition, offre / capacité, lancement, collab 2 et financement se règlent dans Parametres (sections 9, 10, 11, 14, 16)."
LM['A74'] = '— Publicité : budget saisi par année (section 9), plus accompagnement fixe 100 k/mois. Plus de règle en % du CA.'
LM['A88'] = ("PLANCHER DE CAISSE : lire Financement, bloc « Validation 3 ans » — pic de besoin cumulé avant financement, trésorerie minimale après service de la dette, "
             f"couverture OK / INSUFFISANT et financement externe minimum indicatif. Leviers : apport (Parametres!B{ROW['apport']}), mois financés / sécurité "
             f"(B{ROW['fin_mois_dem']}, B{ROW['fin_mois_sec']}), durée de l'emprunt (B{ROW['duree']}), budget publicité (section 9), salaires (section 13).")

# ---------------------------------------------------------------- contrôles
bad = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                for m in ppat.finditer(c.value):
                    if int(m.group(4)) > LASTP: bad.append((ws.title, c.coordinate))
                if 'KPI!B14' in c.value: bad.append((ws.title, c.coordinate, 'ancienne ref'))
assert not bad, bad[:10]
wb.active = wb.sheetnames.index('Dashboard')
wb.save(OUT)
import json
json.dump({'ROW': ROW, 'pv': pv_rows, 'krow': krow, 'arow': arow}, open('/Users/macbookair/Desktop/tairotairo/.excel_work/v3_rows.json', 'w'))
print('saved ; Parametres', LASTP, 'lignes ; Plan ventes flag ligne', FLAGROW, '; KPI', rk, 'lignes')
