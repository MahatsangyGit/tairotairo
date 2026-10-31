# -*- coding: utf-8 -*-
"""Refonte du bloc financement : besoin 6 mois, sécurité calculée, apport 4 M, emprunt auto, validation 3 ans."""
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

PATH = '/Users/macbookair/Desktop/Etude Tairo/Excel base/Modele_Tairo_3_etats.xlsx'
OUT = '/Users/macbookair/Desktop/tairotairo/.excel_work/Modele_Tairo_3_etats_FIN.xlsx'

wb = openpyxl.load_workbook(PATH)
P = wb['Parametres']
F = wb['Financement']

# ------------------------------------------------------------------ styles (identiques au passage précédent)
C_DARK, C_SEC, C_HDRBG, C_INPUT, C_GRAY = '1F4E79', '4472C4', 'D9E1F2', 'FFF6CC', '595959'
F_TITLE = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
F_SEC = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
F_HDR = Font(name='Calibri', size=10, bold=True, color=C_DARK)
FILL_T = PatternFill('solid', fgColor=C_DARK)
FILL_S = PatternFill('solid', fgColor=C_SEC)
FILL_H = PatternFill('solid', fgColor=C_HDRBG)
FILL_I = PatternFill('solid', fgColor=C_INPUT)
B_THIN = Border(bottom=Side(style='thin', color='B2B2B2'))
B_TOP = Border(top=Side(style='thin', color='404040'))
MONEY = '#,##0;(#,##0)'

def font_calc(v, bold=False):
    color = '008000' if (isinstance(v, str) and v.startswith('=') and '!' in v) else '000000'
    return Font(name='Calibri', size=10, bold=bold, color=color)

# ------------------------------------------------------------------ 1) contrôle préalable des refs vers Parametres 114 (sécurité)
ref_pat = re.compile(r"Parametres!\$?([A-H])\$?(\d{1,3})")
refs114 = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                for m in ref_pat.finditer(c.value):
                    if int(m.group(2)) == 114:
                        refs114.append((ws.title, c.coordinate))
assert all(t == 'Financement' for t, _ in refs114), refs114
assert P['A113'].value == 'Apport en fonds propres' and P['A116'].value == "Durée de l'emprunt (mois)"

# ------------------------------------------------------------------ 2) mapping lignes Parametres (insertion de 3 lignes après 116)
INS_AT, K = 117, 3
def pmap(r):
    if r <= 113: return r
    if r == 114: raise ValueError('114 doit être traité à part')
    if r == 115: return 117
    if r == 116: return 118
    return r + K

def remap_ext_param(f):
    return re.sub(r"Parametres!(\$?)([A-H])(\$?)(\d{1,3})",
                  lambda m: f"Parametres!{m.group(1)}{m.group(2)}{m.group(3)}{pmap(int(m.group(4)))}", f)

loc_pat = re.compile(r"(?<![A-Za-z0-9_.!$])(\$?)([A-H])(\$?)(\d{1,3})(?![\d(])")
def remap_local(f):
    parts = re.split(r"([A-Za-z_][A-Za-z0-9_.]*!\$?[A-Z]{1,2}\$?\d{1,3})", f)
    def sub(m):
        r = int(m.group(4))
        if r == 114:
            return f'{m.group(1)}{m.group(2)}{m.group(3)}116'  # plancher embauche (découplé)
        return f'{m.group(1)}{m.group(2)}{m.group(3)}{pmap(r)}'
    return ''.join(p if i % 2 else loc_pat.sub(sub, p) for i, p in enumerate(parts))

# autres onglets (hors Financement, reconstruit plus bas)
for ws in wb.worksheets:
    if ws.title in ('Parametres', 'Financement'):
        continue
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and 'Parametres!' in c.value:
                c.value = remap_ext_param(c.value)

# formules internes de Parametres (avant insertion, sur les anciennes coordonnées)
for row in P.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith('='):
            c.value = remap_local(c.value)

# insertion physique
P.insert_rows(INS_AT, K)
# hauteurs de lignes : réinitialiser puis remettre 18 sur les en-têtes de section (fond 4472C4)
for r in range(INS_AT, P.max_row + 1):
    P.row_dimensions[r].height = None
    fill = P.cell(row=r, column=1).fill
    if fill.fill_type == 'solid' and str(fill.fgColor.rgb).endswith(C_SEC):
        P.row_dimensions[r].height = 18

# ------------------------------------------------------------------ 3) contenu section 14 (lignes 113–119)
SEC14 = [
    (113, 'Apport en fonds propres', 4000000, 'Ar', MONEY,
     "Apport des fondateurs. Le solde du besoin (Financement) est couvert automatiquement par l'emprunt."),
    (114, 'Besoin de démarrage — nombre de mois de charges couverts', 6, 'mois', '0',
     "N mois × (OPEX fixes mensuels du scénario + coûts variables du mois 1). Calcul dans Financement. "
     "Basé sur les hypothèses (pas sur les flux M2–M6) pour éviter une référence circulaire avec le déclencheur d'embauche."),
    (115, "Trésorerie de sécurité — nombre de mois d'OPEX fixes", 3, 'mois', '0',
     "Matelas de sécurité = N mois d'OPEX fixes mensuels (remplace l'ancienne valeur brute de 1 M Ar)."),
    (116, "Plancher de trésorerie pour l'embauche du collaborateur 2", 1000000, 'Ar', MONEY,
     "Découplé de la trésorerie de sécurité. Utilisé uniquement par le déclencheur d'embauche (section 12) et par la validation 3 ans (Financement)."),
    (117, "Taux d'emprunt annuel", 0.09, '%/an', '0.0%', None),
    (118, "Durée de l'emprunt (mois)", 36, 'mois', '0', None),
    (119, 'Référence V2.3 — apport calé à la main pour un plancher de 1 M Ar (documentation)', 36700000, 'Ar', MONEY,
     "Ancien apport itéré manuellement pour que la trésorerie Base ne passe jamais sous 1 M Ar. "
     "Conservé pour comparaison avec le besoin réel calculé (Financement, bloc Validation 3 ans)."),
]
for r, lab, val, unit, fmt, note in SEC14:
    for c in range(1, 9):
        cl = P.cell(row=r, column=c)
        cl.value = None; cl.fill = PatternFill(); cl.border = Border(); cl.number_format = 'General'
    P.cell(row=r, column=1, value=lab).font = Font(name='Calibri', size=10)
    P.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical='top')
    b = P.cell(row=r, column=2, value=val); b.font = Font(name='Calibri', size=10, color='0000FF'); b.fill = FILL_I
    b.number_format = fmt; b.alignment = Alignment(vertical='top')
    u = P.cell(row=r, column=3, value=unit); u.font = Font(name='Calibri', size=10, color=C_GRAY); u.alignment = Alignment(vertical='top')
    d = P.cell(row=r, column=4, value=f'=B{r}'); d.font = font_calc(d.value); d.number_format = fmt; d.alignment = Alignment(vertical='top')
    if note:
        n = P.cell(row=r, column=5, value=note)
        n.font = Font(name='Calibri', size=9, italic=True, color=C_GRAY); n.alignment = Alignment(wrap_text=True, vertical='top')
assert P['D101'].value == '=$B$116', P['D101'].value
P['A101'] = "Trésorerie minimale exigée pour embaucher (plancher, section 14)"

# ------------------------------------------------------------------ 4) reconstruction de Financement
# ancien -> nouveau (lignes référencées par d'autres onglets)
FMAP = {3: 4, 6: 11, 7: 14, 8: 15, 11: 20, 12: 21, 13: 22, 16: 28, 17: 29, 18: 30, 19: 31}
fref = re.compile(r"Financement!(\$?)([A-Z]{1,2})(\$?)(\d{1,3})")
used = set()
for ws in wb.worksheets:
    if ws.title == 'Financement': continue
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and 'Financement!' in c.value:
                used.update(int(m.group(4)) for m in fref.finditer(c.value))
assert used <= set(FMAP), used - set(FMAP)
for ws in wb.worksheets:
    if ws.title == 'Financement': continue
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and 'Financement!' in c.value:
                c.value = fref.sub(lambda m: f"Financement!{m.group(1)}{m.group(2)}{m.group(3)}{FMAP[int(m.group(4))]}", c.value)

# effacement complet
for row in F.iter_rows(min_row=1, max_row=max(F.max_row, 45), max_col=40):
    for c in row:
        c.value = None; c.font = Font(name='Calibri', size=10); c.fill = PatternFill()
        c.border = Border(); c.number_format = 'General'; c.alignment = Alignment()

LASTC = 38  # AL = mois 36
def title(r, text):
    F.cell(row=r, column=1, value=text)
    for c in range(1, LASTC + 1):
        cl = F.cell(row=r, column=c); cl.font = F_TITLE; cl.fill = FILL_T
    F.row_dimensions[r].height = 22
def section(r, text):
    F.cell(row=r, column=1, value=text)
    for c in range(1, LASTC + 1):
        cl = F.cell(row=r, column=c); cl.font = F_SEC; cl.fill = FILL_S
    F.row_dimensions[r].height = 18
def line(r, lab, formula, fmt=MONEY, bold=False, note=None, total=False):
    a = F.cell(row=r, column=1, value=lab); a.font = Font(name='Calibri', size=10, bold=bold)
    b = F.cell(row=r, column=2, value=formula); b.font = font_calc(formula, bold); b.number_format = fmt
    if total:
        a.border = B_TOP; b.border = B_TOP
    if note:
        n = F.cell(row=r, column=3, value=note)
        n.font = Font(name='Calibri', size=9, italic=True, color=C_GRAY); n.alignment = Alignment(vertical='top')

STRESS_CAPEX = "(1+IF(Parametres!$B$140=1,Parametres!$B$144,0))"
title(1, "FINANCEMENT — besoin initial, plan de financement, emprunt et validation sur 3 ans")

section(3, 'BESOIN DE FINANCEMENT INITIAL')
line(4, 'CAPEX (investissements, choc stress inclus)', f"=Parametres!$D$109*{STRESS_CAPEX}")
line(5, 'OPEX fixes mensuels (scénario, année 1)', "=Parametres!$D$75", note='Total OPEX mensuel — Parametres section 9 (multiplicateur scénario et choc stress inclus).')
line(6, 'Coûts variables du mois 1', "=-'P&L'!B7", note='MM, SMS, support, vérification, litiges du premier mois.')
line(7, 'Nombre de mois de démarrage couverts', "=Parametres!$D$114", fmt='0')
line(8, 'Besoin de démarrage (N mois × (OPEX fixes + coûts variables M1))', "=B7*(B5+B6)", bold=True)
line(9, "Nombre de mois d'OPEX fixes en sécurité", "=Parametres!$D$115", fmt='0')
line(10, "Trésorerie de sécurité (N mois × OPEX fixes)", "=B9*B5", bold=True)
line(11, 'TOTAL BESOIN DE FINANCEMENT', "=B4+B8+B10", bold=True, total=True)

section(13, 'PLAN DE FINANCEMENT')
line(14, 'Apport en fonds propres', "=Parametres!$D$113")
line(15, "Emprunt (calcul automatique : besoin non couvert par l'apport)", "=MAX(B11-B14,0)", bold=True)
line(16, 'TOTAL RESSOURCES (= besoin)', "=B14+B15", bold=True, total=True)
line(17, "Part de l'emprunt dans le financement", "=IF(B16=0,0,B15/B16)", fmt='0.0%')

section(19, "EMPRUNT — CONDITIONS")
line(20, 'Montant emprunté', "=B15")
line(21, 'Taux annuel', "=Parametres!$D$117", fmt='0.0%')
line(22, 'Durée (mois)', "=Parametres!$D$118", fmt='0')
line(23, 'Mensualité constante (annuités)', "=IF(B20=0,0,-PMT(B21/12,B22,B20))")
line(24, "Coût total des intérêts sur la durée", "=B29", note="Intérêts → P&L ligne 22 (charges financières) et Trésorerie ligne 9 ; capital → Trésorerie ligne 20 ; dette restante → Bilan ligne 13.")

section(26, "TABLEAU D'AMORTISSEMENT DE L'EMPRUNT (mensuel)")
F.cell(row=27, column=1, value='Mois'); F.cell(row=27, column=2, value='TOTAL')
for i in range(1, 37):
    F.cell(row=27, column=2 + i, value=i).number_format = '0'
for c in range(1, LASTC + 1):
    cl = F.cell(row=27, column=c); cl.font = F_HDR; cl.fill = FILL_H
labels = {28: 'Mensualité', 29: 'Intérêts', 30: 'Capital remboursé', 31: 'Capital restant dû'}
for r, lab in labels.items():
    F.cell(row=r, column=1, value=lab).font = Font(name='Calibri', size=10, bold=(r in (28, 31)))
for r in (28, 29, 30):
    b = F.cell(row=r, column=2, value=f'=SUM(C{r}:AL{r})'); b.font = Font(name='Calibri', size=10, bold=True); b.number_format = MONEY
for i in range(1, 37):
    L = get_column_letter(2 + i)
    F[f'{L}28'] = f'=IF({L}$27>$B$22,0,-PMT($B$21/12,$B$22,$B$20))'
    F[f'{L}29'] = f'=IF({L}$27>$B$22,0,-IPMT($B$21/12,{L}$27,$B$22,$B$20))'
    F[f'{L}30'] = f'=IF({L}$27>$B$22,0,-PPMT($B$21/12,{L}$27,$B$22,$B$20))'
    F[f'{L}31'] = '=$B$20-C30' if i == 1 else f'={get_column_letter(1 + i)}31-{L}30'
    for r in (28, 29, 30, 31):
        F[f'{L}{r}'].number_format = MONEY; F[f'{L}{r}'].font = Font(name='Calibri', size=10, bold=(r == 31))

section(33, "VALIDATION 3 ANS — le financement couvre-t-il la trésorerie sur M0–M36 ?")
line(34, 'Pic de besoin cumulé de trésorerie AVANT financement (flux op. hors intérêts + CAPEX)', "=-MIN(B42:AL42)", bold=True,
     note="Point bas du cumul des flux (ligne 42). C'est le montant qu'il faut financer pour ne jamais être à zéro, hors coût de la dette.")
line(35, 'Financement minimal pour respecter le plancher (pic + plancher, hors service de la dette)', "=B34+Parametres!$D$116")
line(36, 'Trésorerie minimale sur 36 mois APRÈS financement et service de la dette', "=MIN(Tresorerie!D25:AM25)", bold=True)
line(37, 'Couverture du plancher', '=IF(B36>=Parametres!$D$116,"OK — plancher respecté","INSUFFISANT — augmenter apport, emprunt ou durée")', fmt='General', bold=True)
line(38, "Emprunt minimum indicatif pour respecter le plancher (itération suivante, hors intérêts additionnels)", "=B15+MAX(0,Parametres!$D$116-B36)")
line(39, 'Référence V2.3 — apport 36,7 M Ar calé à la main pour un plancher de 1 M', "=Parametres!$D$119")
line(40, 'Écart référence 36,7 M vs financement minimal (positif = référence surdimensionnée)', "=B39-B35")

# ligne de calcul : cumul des flux op. (hors intérêts) + investissement, M0 (col B) .. M36 (col AL)
F.cell(row=42, column=1, value='Cumul flux opérationnels (hors intérêts) + CAPEX, avant financement — M0 (col. B) → M36 (col. AL)').font = Font(name='Calibri', size=9, italic=True, color=C_GRAY)
for i in range(0, 37):
    L = get_column_letter(2 + i)          # B = M0 ... AL = M36
    T = get_column_letter(3 + i)          # Tresorerie : C = M0 ... AM = M36
    flux = f"(SUM(Tresorerie!{T}6:{T}8)+Tresorerie!{T}10+Tresorerie!{T}15)"
    F[f'{L}42'] = f'={flux}' if i == 0 else f'={get_column_letter(1 + i)}42+{flux}'
    F[f'{L}42'].number_format = MONEY; F[f'{L}42'].font = Font(name='Calibri', size=9, color='008000')

F.column_dimensions['A'].width = 78
F.column_dimensions['B'].width = 16
for i in range(3, LASTC + 1):
    F.column_dimensions[get_column_letter(i)].width = 11
F.freeze_panes = None
F.sheet_view.showGridLines = False

# ------------------------------------------------------------------ 5) textes Lisez-moi / Dashboard
LM = wb['Lisez-moi']
LM['A37'] = ("— Financement : apport 4 M Ar (Parametres!B113) ; emprunt calculé automatiquement = besoin − apport, où besoin = CAPEX + 6 mois de charges "
             "(OPEX fixes + coûts variables M1) + 3 mois d'OPEX fixes en sécurité. L'ancien apport 36,7 M (V2.3, calé à la main pour un plancher de 1 M) "
             "est conservé en référence (Parametres!B119) et comparé au besoin réel dans Financement (bloc Validation 3 ans).")
LM['A66'] = ("Fichier distinct de V2.2 / V2.3 (pas une V2.4). Moteur V2.3 (volumes, trigger d'embauche, Ampio+ brut, CAC par face, liquidité KPI, scénarios, stress, IFRS 15) "
             "+ coûts variables de volume V2.2 + plan de financement recalculé (apport 4 M, emprunt automatique).")
LM['A88'] = ("PLANCHER DE CAISSE : lire Financement, bloc « Validation 3 ans » — pic de besoin cumulé avant financement, trésorerie minimale après service de la dette, "
             "couverture OK / INSUFFISANT et emprunt minimum indicatif. Pour rétablir le plancher : augmenter l'apport (Parametres!B113), les mois de démarrage / sécurité "
             "(B114, B115) ou la durée de l'emprunt (B118).")
DD = wb['Dashboard']
assert DD['A33'].value == 'Investissement initial total'
DD['A33'] = 'Besoin de financement initial (CAPEX + démarrage + sécurité)'

# ------------------------------------------------------------------ 6) contrôle final : aucune ref hors bornes
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                for m in ref_pat.finditer(c.value):
                    assert int(m.group(2)) <= P.max_row, (ws.title, c.coordinate, c.value)
wb.save(OUT)
print('saved', OUT, '| Parametres max_row =', P.max_row)
